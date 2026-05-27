# ============================================================
#  AUTOCARES ALEGRE — Hojas de Ruta v3.0
#  Con persistencia en Supabase
# ============================================================
import streamlit as st
import pandas as pd
import pdfplumber
import re, io, json
from supabase import create_client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date, datetime, timedelta

st.set_page_config(page_title="Hojas de Ruta — Autocares Alegre",
                   page_icon="🚌", layout="wide",
                   initial_sidebar_state="collapsed")

CSS = """<style>
header{visibility:hidden;}#MainMenu{visibility:hidden;}
footer{visibility:hidden;}[data-testid="stToolbar"]{display:none;}
[data-testid="stDecoration"]{display:none;}
.block-container{padding-top:0.8rem !important;}
.stTabs [data-baseweb="tab"]{font-size:1rem;font-weight:600;}
</style>"""

COLS_BD = ["nro_hoja","bus_grupo","fecha","cliente","descripcion",
           "plazas","hora_salida","punto_recogida","hora_llegada",
           "destino_regreso","observaciones","conductor","archivo"]

# ── CONEXIONES ───────────────────────────────────────────────
@st.cache_resource
def get_sb():
    return create_client(st.secrets["supabase"]["url"], st.secrets["supabase"]["key"])

def cargar_hojas_ruta(desde, hasta):
    try:
        r = get_sb().table("hojas_ruta").select(
            "id,"+",".join(COLS_BD)+",creado_en"
        ).gte("fecha", str(desde)).lte("fecha", str(hasta)).order("fecha").order("nro_hoja").execute()
        if not r.data: return pd.DataFrame()
        return pd.DataFrame(r.data)
    except Exception as e:
        st.error(f"Error cargando datos: {e}"); return pd.DataFrame()

def guardar_hojas(filas_df):
    """Inserta o actualiza (upsert) por nro_hoja+bus_grupo."""
    sb = get_sb()
    datos = []
    for _, row in filas_df.iterrows():
        d = {}
        for col in COLS_BD:
            val = row.get(col, None)
            if val is None or (isinstance(val, float) and pd.isna(val)) or str(val) in ("","nan","None"):
                d[col] = None
            else:
                d[col] = str(val) if col not in ("fecha",) else str(val)
        datos.append(d)
    try:
        sb.table("hojas_ruta").upsert(datos, on_conflict="nro_hoja,bus_grupo").execute()
        cargar_hojas_ruta.clear()
        return True, len(datos)
    except Exception as e:
        return False, str(e)

def actualizar_fila(id_row, cambios):
    """Actualiza campos de una fila por su id."""
    try:
        get_sb().table("hojas_ruta").update(cambios).eq("id", id_row).execute()
        cargar_hojas_ruta.clear()
        return True
    except Exception as e:
        st.error(f"Error guardando: {e}"); return False

def eliminar_filas(ids):
    try:
        sb = get_sb()
        for i in range(0, len(ids), 100):
            sb.table("hojas_ruta").delete().in_("id", ids[i:i+100]).execute()
        cargar_hojas_ruta.clear()
        return True
    except Exception as e:
        st.error(f"Error eliminando: {e}"); return False

# ── PARSER PDF ───────────────────────────────────────────────
def parse_hora(texto):
    if not texto: return ""
    texto = re.sub(r'_+','',str(texto)).strip().lower()
    texto = re.sub(r'[/].*','', texto)
    texto = texto.replace('am','').replace('pm','').replace('h','').strip().replace('.',':')
    try:
        p = texto.split(':')
        return f"{int(p[0]):02d}:{int(p[1]) if len(p)>1 else 0:02d}"
    except: return ""

def extraer_pdf(pdf_bytes, filename):
    try:
        with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
            grupos_pag = {}; orden = []
            for page in pdf.pages:
                t = page.extract_text() or ""
                m = re.search(r'HOJA DE RUTA\s+\d+\s+(\d+)\s+(\d+)\s+(\d{2}/\d{2}/\d{4})', t)
                if m:
                    nro = m.group(1)
                    if nro not in grupos_pag: grupos_pag[nro]=[]; orden.append(nro)
                    grupos_pag[nro].append(t)
                elif orden: grupos_pag[orden[-1]].append(t)
            if not grupos_pag:
                return [{"error":"No se encontró ninguna HOJA DE RUTA","archivo":filename}]
            return [_parsear_hoja("\n".join(grupos_pag[nro]), nro, filename) for nro in orden]
    except Exception as e:
        return [{"error":str(e),"archivo":filename}]

def _parsear_hoja(texto, nro_hoja, filename):
    try:
        lines = texto.split('\n')
        fecha = ""
        m = re.search(r'HOJA DE RUTA\s+\d+\s+\d+\s+\d+\s+(\d{2}/\d{2}/\d{4})', texto)
        if m:
            partes = m.group(1).split('/')
            fecha = f"{partes[2]}-{partes[1]}-{partes[0]}"
        cliente = ""
        for i, l in enumerate(lines):
            if 'AUTOCARES ALEGRE' in l and i+1 < len(lines):
                nx = lines[i+1].strip()
                if nx and 'TOMAS SANZ' not in nx: cliente = nx; break
        desc_raw = ""
        m_dr = re.search(r'DESCRIPCIÓN\s+CANTIDAD.*?\n(.+?)(?=TIPO\s+IMPORTE)', texto, re.DOTALL)
        if m_dr: desc_raw = m_dr.group(1)
        def es_util(l):
            l=l.strip()
            return (l and not re.match(r'^[\d.,\s€"*]+$',l) and '@' not in l
                    and 'tel:' not in l.lower() and not re.match(r'^PT\d+',l)
                    and 'DNI' not in l and 'Teléfono' not in l
                    and 'CORREO' not in l.upper() and 'A CTA' not in l
                    and 'PAGADO' not in l and not l.startswith('http')
                    and 'maps.app' not in l and not re.match(r'^\d+\s*€',l)
                    and 'share.google' not in l)
        desc_lines = [l.strip() for l in desc_raw.split('\n') if es_util(l.strip())]
        desc = " / ".join(desc_lines[:2]) if desc_lines else ""
        fuente = desc_raw or texto
        # DESTINO
        destino_global = ""
        for pat in [
            r'SALON DESTINO[:\s]+([^\n•\(]+)',
            r'MASIA[:\s]+([^\n•\(]+)',
            r'[Cc]ampamento[:\s]+([^\n•\(,]+)',
            r'[Dd]estino(?:\s+final)?[:\s]+([A-ZÁÉÍÓÚa-záéíóúñÑ][^\n•]{2,55})',
            r'[Ll]ugar[:\s]+([^\n\(•]{3,55})',
        ]:
            m = re.search(pat, fuente, re.IGNORECASE)
            if m:
                d = m.group(1).strip().split('(')[0].strip()
                if not d.startswith('http') and len(d) > 3:
                    destino_global = d[:60]; break
        if not destino_global:
            m = re.search(r'^([A-ZÁÉÍÓÚa-záéíóúñÑ][A-ZÁÉÍÓÚa-záéíóúñÑ\s]{2,25})\s+[Ss]alida\s+\d', fuente, re.MULTILINE)
            if m: destino_global = m.group(1).strip()
        if not destino_global:
            for i, l in enumerate(lines):
                l2 = l.strip().rstrip(',.')
                if (re.match(r'^[A-ZÁÉÍÓÚ][A-ZÁÉÍÓÚa-záéíóúñÑ\s]{3,25}$', l2) and
                    i+1<len(lines) and re.search(r'salida|llegada|horario', lines[i+1], re.IGNORECASE)):
                    destino_global = l2; break
        # HORAS
        h_sal = h_reg = ""
        pats_sal = [
            (r'[Ss]alida[:\s]+(\d{1,2}[:.]\d{2})[hH]?\s+[Ll]legada[:\s]+(\d{1,2}[:.]\d{2})', True),
            (r'[Hh]ora[s]?\s*salida[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Hh]orario\s+de\s+salida[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Hh]ora\s+de\s+(?:salida|recogida)[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Hh]\.\s*[Ss]alida[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'SALIDA[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Ee]ixida[:\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Rr]ecogida[^.]+?a las (\d{1,2}[:.]\d{2})', False),
            (r'^[Ss]alida[":\s]+(\d{1,2}[:.]\d{2})', False),
            (r'[Ss]alida(?:[^\n])*?(\d{1,2}[:.]\d{2})\s*h\b', False),
        ]
        for pat, doble in pats_sal:
            if h_sal: break
            flags = re.MULTILINE if pat.startswith('^') else 0
            m = re.search(pat, fuente, flags)
            if m:
                h_sal = parse_hora(m.group(1))
                if doble and len(m.groups())>1 and m.group(2): h_reg = parse_hora(m.group(2))
        pats_reg = [
            r'[Hh]rs?\s*regreso[:\s]+(\d{1,2}[:.]\d{2})',
            r'[Hh]orario\s+de\s+regreso[:\s]+(\d{1,2}[:.]\d{2})',
            r'[Hh]ora\s+de\s+regreso[^\d]+(\d{1,2}[:.]\d{2})',
            r'[Hh]\.\s*[Rr]egreso[^\d]+(\d{1,2}[:.]\d{2})',
            r'LLEGADA[:\s]+(\d{1,2}[:.]\d{2})',
            r'^[Rr]egreso[":\s]+(\d{1,2}[:.]\d{2})',
            r'[Tt]ornada[^\d]+(\d{1,2}[:.]\d{2})',
            r'[Ss]alida de [^:]+:\s*(\d{1,2}[:.]\d{2})',
            r'[Vv]uelta[^\d]*?(\d{1,2}[:.]\d{2})',
            r'[Ll]legada\s+al\s+colegio[:\s]+(\d{1,2}[:.]\d{2})',
        ]
        if not h_reg:
            for pat in pats_reg:
                flags = re.MULTILINE if pat.startswith('^') else 0
                m = re.search(pat, fuente, flags)
                if m: h_reg = parse_hora(m.group(1)); break
        # PLAZAS
        plazas_global = ""
        m_alu = re.search(r'[Aa]lumnos[:\s]+(\d+)', fuente)
        m_pro = re.search(r'[Pp]rofesores[:\s]+(\d+)', fuente)
        if m_alu:
            total = int(m_alu.group(1)) + (int(m_pro.group(1)) if m_pro else 0)
            plazas_global = str(total)
        else:
            for pat in [r'[Nn]úmero\s+de\s+plazas[^:\d]*[:\s]+(\d+)',
                        r'[Pp]lazas[:\s]+(\d+)', r'(\d+)\s+plazas?\b',
                        r'(\d+)\s+(?:[Pp]asajeros|[Pp]ersonas|[Pp]laces?)\b',
                        r'(\d+)\s*(?:PAX|pax|Pax)\b', r'(\d+)pax\b',
                        r'(\d+)\s+seater', r'(\d+)\s+persones']:
                m = re.search(pat, fuente)
                if m and 1 < int(m.group(1)) < 500:
                    plazas_global = m.group(1); break
        # OBSERVACIONES estructuradas
        partes_obs = []
        if destino_global: partes_obs.append(f"Destino: {destino_global}")
        if h_sal:          partes_obs.append(f"Salida: {h_sal}")
        if h_reg:          partes_obs.append(f"Regreso: {h_reg}")
        if plazas_global:  partes_obs.append(f"{plazas_global} plazas")
        for dl in desc_lines[:3]:
            if (dl not in destino_global and not re.search(r'\d{1,2}[:.]\d{2}',dl)
                    and 'plazas' not in dl.lower() and 'pax' not in dl.lower() and len(dl)>5):
                partes_obs.append(dl); break
        obs = " | ".join(partes_obs)[:180]
        # GRUPOS/BUSES
        grupos = []
        buses_check = re.findall(r'(?:^|[•\n\s])BUS\s+(\d{1,2})\s*[:(]', texto, re.IGNORECASE|re.MULTILINE)
        if buses_check:
            ida_sec = re.search(r'(?:AUTOBUSES DE )?IDA.*?(?=(?:AUTOBUSES DE )?REGRESO|$)', texto, re.DOTALL|re.IGNORECASE)
            reg_sec = re.search(r'(?:AUTOBUSES DE )?REGRESO.*', texto, re.DOTALL|re.IGNORECASE)
            buses_ida = re.findall(r'BUS\s+(\d+).*?\(Plazas[:\s]+_*(\d*)\D*\).*?Hora[:\s]+([0-9_.:\-hHaAmM]+).*?recogida[:\s]+([^\n•]+)',
                                   ida_sec.group(0) if ida_sec else "", re.IGNORECASE)
            buses_reg = re.findall(r'BUS\s+(\d+).*?\(Plazas[:\s]+_*(\d*)\D*\).*?Hora[:\s]+([0-9_.:\-hHaAmM]+).*?[Dd]estino[:\s]+([^\n•]+)',
                                   reg_sec.group(0) if reg_sec else "", re.IGNORECASE)
            def limpiar(s): return re.sub(r'_+','',s).strip().strip('_ ')
            bd = {}
            for num,plazas,hora,recogida in buses_ida:
                n=int(num)
                if n not in bd: bd[n]={"grupo":str(n),"plazas":plazas,"hora_salida":"","hora_llegada":"","destino":"","destino_regreso":"","obs":obs}
                bd[n]["hora_salida"]=parse_hora(hora); bd[n]["plazas"]=plazas; bd[n]["destino"]=limpiar(recogida)[:60]
            for num,plazas,hora,dest_reg in buses_reg:
                n=int(num)
                if n not in bd: bd[n]={"grupo":str(n),"plazas":plazas,"hora_salida":"","hora_llegada":"","destino":"","destino_regreso":"","obs":obs}
                bd[n]["hora_llegada"]=parse_hora(hora); bd[n]["destino_regreso"]=limpiar(dest_reg)[:60]
                if not bd[n]["plazas"]: bd[n]["plazas"]=plazas
            grupos = [bd[k] for k in sorted(bd)]
        if not grupos:
            grupos=[{"grupo":"1","plazas":plazas_global,"hora_salida":h_sal,
                     "hora_llegada":h_reg,"destino":destino_global,"destino_regreso":"","obs":obs}]
        return {"archivo":filename,"nº_hoja":nro_hoja,"fecha":fecha,
                "cliente":cliente,"descripcion":desc,"grupos":grupos}
    except Exception as e:
        return {"error":str(e),"archivo":filename,"nº_hoja":nro_hoja}

def hojas_a_df(hojas_anidadas):
    hojas = []
    for item in hojas_anidadas:
        if isinstance(item, list): hojas.extend(item)
        else: hojas.append(item)
    filas = []
    for h in hojas:
        if "error" in h:
            filas.append({"nro_hoja":"ERROR","bus_grupo":"1","fecha":"","cliente":"",
                          "descripcion":"","plazas":"","hora_salida":"","punto_recogida":"",
                          "hora_llegada":"","destino_regreso":"","conductor":"",
                          "observaciones":h["error"][:80],"archivo":h.get("archivo","")}); continue
        for g in h.get("grupos",[]):
            filas.append({
                "nro_hoja":        h.get("nº_hoja",""),
                "bus_grupo":       g.get("grupo","1"),
                "fecha":           h.get("fecha",""),
                "cliente":         h.get("cliente",""),
                "descripcion":     h.get("descripcion",""),
                "plazas":          g.get("plazas",""),
                "hora_salida":     g.get("hora_salida",""),
                "punto_recogida":  g.get("destino",""),
                "hora_llegada":    g.get("hora_llegada",""),
                "destino_regreso": g.get("destino_regreso",""),
                "observaciones":   g.get("obs",""),
                "conductor":       "",
                "archivo":         h.get("archivo",""),
            })
    return pd.DataFrame(filas)

# ── VERIFICACIÓN CON CONDUCTORES ─────────────────────────────
@st.cache_data(ttl=30)
def cargar_registros_conductores():
    try:
        r = get_sb().table("datos_brutos").select(
            "conductor,fecha,hoja_servicio,servicios_horas,tipo_registro"
        ).execute()
        if not r.data: return pd.DataFrame()
        df = pd.DataFrame(r.data)
        df["hoja_servicio"] = df["hoja_servicio"].fillna("").astype(str).str.strip()
        return df
    except: return pd.DataFrame()

def primera_hora(raw):
    try:
        items = json.loads(str(raw or "[]"))
        if items:
            ini = items[0].get("inicio","").strip()
            fin = items[0].get("fin","").strip()
            return (ini if ini!="00:00" else ""), (fin if fin!="00:00" else "")
    except: pass
    return "",""

def verificar_conductores(df_hojas, df_cond):
    """Cruza hojas con registros de conductores. Devuelve df con columnas de verificación."""
    resultados = []
    for _, row in df_hojas.iterrows():
        nro = str(row.get("nro_hoja","")).strip().lstrip("0")
        matches = pd.DataFrame()
        if not df_cond.empty and nro:
            mask = df_cond["hoja_servicio"].str.lstrip("0") == nro
            matches = df_cond[mask]
        if matches.empty:
            resultados.append({**row.to_dict(), "conductor_verificado":"",
                               "estado_verif":"⚠️ Sin registro"})
        else:
            for _, reg in matches.iterrows():
                conductor = reg.get("conductor","")
                tipo_r = str(reg.get("tipo_registro","extra") or "extra")
                h_ini_c, h_fin_c = primera_hora(reg.get("servicios_horas",""))
                if tipo_r == "jornada":
                    estado = "📋 Jornada"
                else:
                    h_sal = str(row.get("hora_salida","")).strip()
                    h_reg = str(row.get("hora_llegada","")).strip()
                    def cmp(a,b):
                        if not a or not b: return "sin_dato"
                        try:
                            return "ok" if datetime.strptime(a,"%H:%M").time()==datetime.strptime(b,"%H:%M").time() else "dif"
                        except: return "sin_dato"
                    es = cmp(h_sal, h_ini_c); el = cmp(h_reg, h_fin_c)
                    if es=="ok" and el=="ok": estado="✅ Coincide"
                    elif es=="sin_dato" or el=="sin_dato": estado="⚠️ Verificar horario"
                    else:
                        difs=[]
                        if es=="dif": difs.append(f"Sal: PDF {h_sal}/Cond. {h_ini_c}")
                        if el=="dif": difs.append(f"Reg: PDF {h_reg}/Cond. {h_fin_c}")
                        estado = "❌ "+" | ".join(difs)
                resultados.append({**row.to_dict(),
                                    "conductor_verificado": conductor,
                                    "estado_verif": estado})
    return pd.DataFrame(resultados)

# ── EXCEL ────────────────────────────────────────────────────
def generar_excel(df):
    output = io.BytesIO()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title="Hojas de Ruta"
    s = Side(style="thin",color="CCCCCC")
    borde = Border(left=s,right=s,top=s,bottom=s)
    for cn,col in enumerate(df.columns,1):
        c=ws.cell(row=1,column=cn,value=col)
        c.font=Font(bold=True,color="FFFFFF",name="Calibri",size=10)
        c.fill=PatternFill("solid",fgColor="1A3A5C")
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        c.border=borde
    ws.row_dimensions[1].height=24
    for rn,(_, row) in enumerate(df.iterrows(),2):
        est=str(row.get("estado_verif",""))
        bg="C6EFCE" if "✅" in est else "FFC7CE" if "❌" in est else "FFEB9C" if "⚠️" in est else "FFFFFF"
        fill=PatternFill("solid",fgColor=bg)
        has_estado="estado_verif" in df.columns
        for cn,val in enumerate(row,1):
            c=ws.cell(row=rn,column=cn,value=str(val) if val is not None else "")
            c.font=Font(name="Calibri",size=9)
            c.alignment=Alignment(horizontal="left",vertical="center",wrap_text=True)
            c.border=borde
            if has_estado and cn==df.columns.get_loc("estado_verif")+1:
                c.fill=fill; c.font=Font(name="Calibri",size=9,bold=True)
        ws.row_dimensions[rn].height=15
    anchos={"nro_hoja":10,"bus_grupo":7,"fecha":12,"cliente":28,"descripcion":30,
            "plazas":8,"hora_salida":12,"punto_recogida":28,"hora_llegada":12,
            "destino_regreso":28,"observaciones":45,"conductor":22,"archivo":18,
            "conductor_verificado":22,"estado_verif":35}
    for cn,col in enumerate(df.columns,1):
        ws.column_dimensions[get_column_letter(cn)].width=anchos.get(col,14)
    ws.freeze_panes="A2"
    wb.save(output); output.seek(0); return output.getvalue()

# ── INTERFAZ ─────────────────────────────────────────────────
def main():
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown("<h2 style='color:#1a3a5c;margin-bottom:0;'>🚌 Hojas de Ruta — Autocares Alegre</h2>",
                unsafe_allow_html=True)

    tab1, tab2, tab3 = st.tabs(["📄 Extraer PDFs", "📋 Servicios guardados", "🔍 Verificar conductores"])

    # ══════════════════════════════════════════════════════════
    # TAB 1 — Extraer PDFs y previsualizar antes de guardar
    # ══════════════════════════════════════════════════════════
    with tab1:
        st.markdown("### Subir y extraer PDFs")
        st.caption("Extrae los datos, revisa y corrige si es necesario, luego guarda en la base de datos.")

        if "uploader_key" not in st.session_state: st.session_state["uploader_key"]=0
        col_up1, col_up2 = st.columns([4,1])
        with col_up1:
            archivos = st.file_uploader("Selecciona PDFs (puedes subir varios a la vez con Ctrl+clic)",
                                        type=["pdf"], accept_multiple_files=True,
                                        key=f"up_{st.session_state['uploader_key']}")
        with col_up2:
            st.markdown("<br>",unsafe_allow_html=True)
            if st.button("🗑️ Limpiar", use_container_width=True):
                st.session_state["uploader_key"]+=1
                st.session_state.pop("df_preview",None)
                st.rerun()

        if archivos:
            st.markdown(f"**{len(archivos)} PDF(s) cargados**")
            if st.button(f"📄 Extraer datos", type="primary", use_container_width=True):
                hojas=[]
                barra=st.progress(0, text="Extrayendo...")
                for i,archivo in enumerate(archivos):
                    barra.progress((i+1)/len(archivos), text=f"Leyendo {archivo.name}...")
                    hojas.extend(extraer_pdf(archivo.read(), archivo.name))
                barra.empty()
                df = hojas_a_df(hojas)
                st.session_state["df_preview"] = df
                st.rerun()

        if "df_preview" in st.session_state:
            df_prev = st.session_state["df_preview"]
            st.success(f"✅ {len(df_prev)} servicio(s) extraídos — **revisa y corrige antes de guardar**")

            df_ed = st.data_editor(df_prev, use_container_width=True, hide_index=True,
                column_config={
                    "nro_hoja":        st.column_config.TextColumn("Nº Hoja",     disabled=True, width="small"),
                    "bus_grupo":       st.column_config.TextColumn("Bus",         disabled=True, width="small"),
                    "fecha":           st.column_config.TextColumn("Fecha",       disabled=True, width="small"),
                    "cliente":         st.column_config.TextColumn("Cliente",     disabled=True, width="medium"),
                    "descripcion":     st.column_config.TextColumn("Descripción", disabled=True, width="large"),
                    "plazas":          st.column_config.TextColumn("Plazas",      width="small"),
                    "hora_salida":     st.column_config.TextColumn("H. Salida",   width="small"),
                    "punto_recogida":  st.column_config.TextColumn("Punto/Destino", width="medium"),
                    "hora_llegada":    st.column_config.TextColumn("H. Llegada",  width="small"),
                    "destino_regreso": st.column_config.TextColumn("Dest. Regreso", width="medium"),
                    "observaciones":   st.column_config.TextColumn("Observaciones", width="large"),
                    "conductor":       st.column_config.TextColumn("Conductor",   width="medium"),
                    "archivo":         st.column_config.TextColumn("Archivo",     disabled=True, width="small"),
                }, key="editor_preview")

            st.session_state["df_preview"] = df_ed

            col_g1, col_g2 = st.columns(2)
            with col_g1:
                if st.button("💾 Guardar en base de datos", type="primary", use_container_width=True):
                    with st.spinner("Guardando..."):
                        ok, res = guardar_hojas(df_ed)
                    if ok:
                        st.success(f"✅ {res} servicio(s) guardados. Si ya existían, se han actualizado.")
                        st.session_state.pop("df_preview",None)
                        st.rerun()
                    else:
                        st.error(f"Error: {res}")
            with col_g2:
                st.download_button("⬇️ Excel de esta extracción",
                    data=generar_excel(df_ed),
                    file_name=f"extraccion_{date.today().strftime('%Y_%m_%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True)

    # ══════════════════════════════════════════════════════════
    # TAB 2 — Tabla acumulada (Supabase)
    # ══════════════════════════════════════════════════════════
    with tab2:
        st.markdown("### Servicios guardados")

        # Filtro de fechas
        col_f1, col_f2, col_f3, col_f4 = st.columns([2,2,2,1])
        hoy = date.today()
        primer_dia = hoy.replace(day=1)
        with col_f1:
            desde = st.date_input("Desde", value=primer_dia, format="DD/MM/YYYY", key="tab2_desde")
        with col_f2:
            hasta = st.date_input("Hasta", value=hoy, format="DD/MM/YYYY", key="tab2_hasta")
        with col_f3:
            st.markdown("<br>",unsafe_allow_html=True)
            filtro_cli = st.text_input("Filtrar cliente", placeholder="Escribe parte del nombre...", key="tab2_cli")
        with col_f4:
            st.markdown("<br>",unsafe_allow_html=True)
            if st.button("↺ Actualizar", use_container_width=True, key="tab2_upd"):
                cargar_hojas_ruta.clear(); st.rerun()

        df_bd = cargar_hojas_ruta(desde, hasta)

        if df_bd.empty:
            st.info("No hay servicios guardados en ese periodo.")
        else:
            if filtro_cli:
                df_bd = df_bd[df_bd["cliente"].fillna("").str.lower().str.contains(filtro_cli.lower())]

            st.caption(f"**{len(df_bd)} servicios** en el periodo seleccionado. Edita directamente y pulsa **Guardar cambios**.")

            # Mostrar sin columnas técnicas
            cols_mostrar = ["nro_hoja","bus_grupo","fecha","cliente","descripcion",
                           "plazas","hora_salida","punto_recogida","hora_llegada",
                           "destino_regreso","observaciones","conductor"]
            df_show = df_bd[["🗑️"]+cols_mostrar if "🗑️" in df_bd.columns else cols_mostrar].copy()
            df_show.insert(0, "🗑️", False)

            df_ed2 = st.data_editor(df_show, use_container_width=True, hide_index=True,
                column_config={
                    "🗑️":              st.column_config.CheckboxColumn("🗑️",default=False,width="small"),
                    "nro_hoja":        st.column_config.TextColumn("Nº Hoja",     disabled=True, width="small"),
                    "bus_grupo":       st.column_config.TextColumn("Bus",         disabled=True, width="small"),
                    "fecha":           st.column_config.TextColumn("Fecha",       disabled=True, width="small"),
                    "cliente":         st.column_config.TextColumn("Cliente",     disabled=True, width="medium"),
                    "descripcion":     st.column_config.TextColumn("Descripción", disabled=True, width="large"),
                    "plazas":          st.column_config.TextColumn("Plazas",      width="small"),
                    "hora_salida":     st.column_config.TextColumn("H. Salida",   width="small"),
                    "punto_recogida":  st.column_config.TextColumn("Punto/Destino", width="medium"),
                    "hora_llegada":    st.column_config.TextColumn("H. Llegada",  width="small"),
                    "destino_regreso": st.column_config.TextColumn("Dest. Regreso", width="medium"),
                    "observaciones":   st.column_config.TextColumn("Observaciones", width="large"),
                    "conductor":       st.column_config.TextColumn("Conductor",   width="medium"),
                }, key="editor_bd")

            col_s1, col_s2, col_s3 = st.columns([2,2,2])
            with col_s1:
                if st.button("💾 Guardar cambios", type="primary", use_container_width=True):
                    # Actualizar fila por fila las que hayan cambiado
                    cols_edit = ["plazas","hora_salida","punto_recogida","hora_llegada",
                                 "destino_regreso","observaciones","conductor"]
                    actualizados = 0
                    for i, (_, row_ed) in enumerate(df_ed2.iterrows()):
                        if i >= len(df_bd): break
                        id_row = int(df_bd.iloc[i]["id"])
                        cambios = {}
                        for col in cols_edit:
                            val_nuevo = str(row_ed.get(col,"") or "").strip()
                            val_viejo = str(df_bd.iloc[i].get(col,"") or "").strip()
                            if val_nuevo != val_viejo:
                                cambios[col] = val_nuevo or None
                        if cambios:
                            actualizar_fila(id_row, cambios)
                            actualizados += 1
                    if actualizados:
                        st.success(f"✅ {actualizados} fila(s) actualizadas.")
                        cargar_hojas_ruta.clear(); st.rerun()
                    else:
                        st.info("No hay cambios que guardar.")

            with col_s2:
                ids_borrar = df_bd.iloc[df_ed2[df_ed2["🗑️"]==True].index]["id"].tolist() if "🗑️" in df_ed2.columns and df_ed2["🗑️"].any() else []
                if ids_borrar:
                    if st.button(f"🗑️ Eliminar {len(ids_borrar)} fila(s)", type="secondary", use_container_width=True):
                        if eliminar_filas([int(i) for i in ids_borrar]):
                            st.success("Eliminado."); st.rerun()
                else:
                    st.button("🗑️ Eliminar seleccionadas", disabled=True, use_container_width=True)

            with col_s3:
                st.download_button("⬇️ Excel del periodo",
                    data=generar_excel(df_bd[cols_mostrar]),
                    file_name=f"hojas_{desde}_{hasta}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True)

    # ══════════════════════════════════════════════════════════
    # TAB 3 — Verificar con conductores (re-ejecutable)
    # ══════════════════════════════════════════════════════════
    with tab3:
        st.markdown("### Verificar con registros de conductores")
        st.caption("Cruza los servicios guardados con lo que han registrado los conductores. Puedes ejecutarlo varias veces.")

        col_v1, col_v2 = st.columns([2,2])
        with col_v1:
            desde_v = st.date_input("Desde", value=date.today().replace(day=1),
                                    format="DD/MM/YYYY", key="tab3_desde")
        with col_v2:
            hasta_v = st.date_input("Hasta", value=date.today(),
                                    format="DD/MM/YYYY", key="tab3_hasta")

        if st.button("🔍 Verificar ahora", type="primary", use_container_width=True):
            with st.spinner("Cargando datos..."):
                df_hojas = cargar_hojas_ruta(desde_v, hasta_v)
                df_cond  = cargar_registros_conductores()
            if df_hojas.empty:
                st.warning("No hay servicios guardados en ese periodo.")
            else:
                df_verif = verificar_conductores(df_hojas, df_cond)
                st.session_state["df_verif"] = df_verif
                cargar_registros_conductores.clear()
            st.rerun()

        if "df_verif" in st.session_state:
            df_verif = st.session_state["df_verif"]
            total   = len(df_verif)
            n_ok    = df_verif["estado_verif"].str.contains("✅").sum()
            n_jorn  = df_verif["estado_verif"].str.contains("📋").sum()
            n_warn  = df_verif["estado_verif"].str.contains("⚠️").sum()
            n_err   = df_verif["estado_verif"].str.contains("❌").sum()
            m1,m2,m3,m4,m5 = st.columns(5)
            m1.metric("Total",          total)
            m2.metric("✅ Coincide",    int(n_ok))
            m3.metric("📋 Jornada",     int(n_jorn))
            m4.metric("⚠️ Revisar",     int(n_warn))
            m5.metric("❌ Diferencia",  int(n_err))

            filtro = st.radio("Mostrar:", ["Todos","Solo ⚠️ y ❌","Solo ❌"],
                              horizontal=True, key="tab3_filtro")
            df_v = df_verif
            if filtro=="Solo ⚠️ y ❌": df_v=df_verif[df_verif["estado_verif"].str.contains("⚠️|❌")]
            elif filtro=="Solo ❌":    df_v=df_verif[df_verif["estado_verif"].str.contains("❌")]

            cols_v = ["nro_hoja","fecha","cliente","plazas","hora_salida","hora_llegada",
                      "punto_recogida","conductor","conductor_verificado","estado_verif"]
            st.dataframe(df_v[[c for c in cols_v if c in df_v.columns]],
                         use_container_width=True, hide_index=True)

            col_d1, col_d2 = st.columns(2)
            with col_d1:
                st.download_button("⬇️ Excel verificación (filtrado)",
                    data=generar_excel(df_v),
                    file_name=f"verif_{date.today().strftime('%Y_%m_%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary", use_container_width=True)
            with col_d2:
                st.download_button("⬇️ Excel verificación (completo)",
                    data=generar_excel(df_verif),
                    file_name=f"verif_completo_{date.today().strftime('%Y_%m_%d')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True)

if __name__ == "__main__":
    main()
