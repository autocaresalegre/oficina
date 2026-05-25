# ============================================================
#  AUTOCARES ALEGRE v9.0
# ============================================================
import streamlit as st
import pandas as pd
import json
from supabase import create_client
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fpdf import FPDF
from datetime import date, datetime, time
import io

st.set_page_config(page_title="Autocares Alegre", page_icon="🚌",
                   layout="centered", initial_sidebar_state="collapsed")

LOGO_B64 = "/9j/4AAQSkZJRgABAQAAAQABAAD/2wCEAAMCAggKCAgJCAgICwgICAgICAkKCgoKCggKCwgICAgICAgICAgKCAgICggICAoKCAgICgkKCAoLDQoIDQgICggBAwQEBgUGCgYGCg8OCg0QDw8PDw4NDQ0ODQ8PDRAQDhAPEA0QDw0NEA4NEA8NDQ4PDQ0QDQ0NDQ4NDQ0ODQ0NDf/AABEIAGQAyAMBIgACEQEDEQH/xAAdAAACAwADAQEAAAAAAAAAAAAABwEICQQFBgMC/8QAVxAAAgIBAgMFBAMGDREJAQAAAQIDBBEABQYSIQcIEzFBFFFhcQkigRUjMlKRsSRCU1RVYnKCkpShwfAWGSUzQ0WDhJOVorK0tdHU1Rc1VmR1wsPT4Rj/xAAcAQEAAQUBAQAAAAAAAAAAAAAABQEDBAYHAgj/xAA0EQABBAECAwYFAQkBAAAAAAABAAIDEQQFIRIxUQYiQWGhsRNxgcHRghQyUpGSouHw8RX/2gAMAwEAAhEDEQA/ANU9GjRoiNGjRoiNGjUZ0RTo1GjOqIp0ajOp0tEaNRnRnVUU6NRnQNEU6NQToB1S0U6NGjVURo0aNURGvjPr6k6/Eo0VDyVF+8F3k7c9qerTleCpA7wloziSw6nldjJgMkYYFVVCuQOYlsqFRcW/ThudZ5xJ+OJpQ/z5w3N/KdWC7WO6JuAtTS0Ak8E0jyiMuscsRdizKechXUMTggqQMDBwTpa793eN6gjaWWhJyKCWKPHIVHqSiOWIHryg67Vpc2nMhYyNzBYogncnzXy/rmJrMuTJLkRyGiaIvhA8OGuQpOzuu94qeSwm37hIZDLk1J2xz8wDO0MpH4eQuY2PXpynOV1OqwcK7sYrVWZScxWYJc59FkViM/Fcg/AnUa1vWdB45+PGbQI3A5Wtv7N9sDFi/Cy3W5poE86WrejRo1zpdyRo0aNEUE6S/eB72Gz7IoF2cvbZOeGjX5ZLUg64cxlkWGIkFRLYeKLIwCxwNeC76vfGj2WD2Sp4cu8WY+aGNhzJTiJZfbLSggkZVhBFkGd1bqEjmdMmd83+exNJYtTSTWJnaSWaVi8kjEnLFz5D0CpyoigIoVVVRLYeAZhxv/dRXp/rq16W5Xjr7NWStLarQHxbEjzsks0cRwsUSxpKOfp9aZWPT10w+913+bmz7qdsoUqsrQ14pbE1l5cB5QzRxJFCUP1ECuxZ+viKAPMihHdu4aFriLY65zh9zrSnH/l29tIPQ/Vb2flPTyPp5juu+NvBl4q3+UnK+2rGp/awU61bA+Rhb7c6lv2OAyhobyaTW++4H5Rat903thsbzsVPc7UMUU9h7qOkXN4f3i9ZqqyeIzMA4hDEMzYJODjGnCTpO90Dhg1eGNjhYYcbdBJKPdJMvtEvoP7pKxGRnHnk9dOHOtXnrjcG8rNIqEd77v8Am57XvEm3bTBt5WrHEbMtqOeYvLIgk8ONILVVY1iRk5ixlLs/lHyEyI4/ShcVfqex/wAStf8AVdJnvO8RCzxJvs6+R3OzD/FitLp8P0Ofn0PrpUybtECQZYwR5guoI+YJ1t8GHF8NoLATQvZFbd/pO+Kz6bKPlSsfz7nrSHu29pE+47Dte4WliWxdpxzzCJWWLmOQxjV2dlVscwBdsZ8zrCRt7h9JEJ9FVgzN+1VQcszeSgdSSB660I7yHbRPs3CfD/DtZzFuVjZaS33Rir1a61UjnVHUq0c1qbxI0bCnw4p2BVgucPNwmngbG0Ak+lbomf3gfpJ9vozSVdqh9vtRMUklEgjpRuGw8ZnTxZJ5E68whiMYZSplByBVvdfpMOLHYlH2qFcnCx05GIHoOee5Lk/HlGfxR5CrMcYAAAAAAAA6AAeQAHoPTp019dvrvLzCBJJWXoyxI0pB9xWJWbPw5dZseBBGNxfW90VgL/f74sfP9lOTP6nXrj/WibH8uuu//uTi39mp/wDIU8f7L/PpUQ9n+5tjl2rdmz7tuun81Y6/UnZ3uY/C2ndx89uvD89Ua9/Bg/halp00PpAuLEx/ZJHI/Va0DZ+xEj/ONXJ7lnfofd5jtu6pDHuYRpK8sIZILqL+Gojd5Ghsxg8zJ4jrIgMicvLLHFl3f2SxEMzVbcIyBzT1p4VyfJQ08Ualj7s5PoDrm8FcaPRu078bMrULUFvmX8LkjcNMg8v7dD4kJ9CsjA9G1amw4pG00AHyRa+d9jvBWtj2dLdFa7W579apALCO8R5kmnlLJFLC5xDXlIxKuDg6o630ofFX6lsQ+VO3/Pu2nH9K/wAQq1PYYVbrJbnugZ6MqVGhB+OPagP335c47F1Fxzuq58uYgZ+WSNYmn4rHQguaCTfuito30nnFf4uyD/ErP/Uzqxvcc73W9b3fvVtyWgYYKiTq1aCSFldpfD5X8S1YDIw5iOiHKnqfTLj7tw/qsf8ADX/jq8f0bM6w0uKtzVgViq1q6sOo50S1OVUjpkGVCQGyOZQffrJysOMR91gDiQBt4k7KxO8RRukdyAJPyAXW7hCj3JVhA5HtuIgPIK05EYH2EAY0a53ZHsxl3LbYh626+c+oRxI3+ih0a3XVNVGC5kQF937r5o0LQv8A1GSTHbvfa1qRo0agnXGl9PIOlV3lO3Wvsm1T3puVpMGGlXLcpt2mVzBApAYqp5S8rhW8KJHkIxGdNOR8DJ8h1J9326xq77veFbed5kELt9zdtL1aS5PLK6sVs3SpVesrL4URPNiGMMrDx3Gs3Dx/jvo8huUST414zs3LNi7emMlmzI8s0jdACTnkRST4cMIHhxx82I41VepQszN4g7q26VOH5N8vha0Pi1I61SQMLVhbE0cKzOn1BWCiTxBFIGlIU8wh5cOyO4J2ZbNNcfct5t1ETbpkSpUsSRp4tkRwzpbdJGHPHW5h4SkFfaCz/hV0Ic30mvbJSn2zb6VK5Xnaa+bM6xSrJyxwwS8rN4ZblBlkQDPLnr1zrYZMlwlbDGKHiUSA+ju2YS8W0yVytWpuFw/teWOOspPyNofl1XntI3prdjcp4zzm9auzw482E88kkCj48siIPlq0/wBH4Fg/qo3J8AUNiZY5PxTIZpZBn4+ywH96NV17vXBvtG77FTODz7jt4YfjLDKliVfkY4JM/PWQHVLI7+EAe5RbwcPbSIYIIR5QwxxD94ip/NrkbhbCI7t+CiM5+SqWP8g1ydLXvKcWex8P71aBw1fbLsifuxA4jA+LMVUfE60to4nAdSiw7KtctHkY8+43TyP689y0SrD4884I1u1w52S7ZBBFBFt9NY4Y0jRRBF0CqFHXkyT06knJOsGNucx+GYyVaIo0bDoUZCCjKR5MhUMPcRpuHvg8Wf8AiHcP4NT8/sf59bdm4r5+ERuAA6308kW0EXBlJSCKlUEHORDGMY6+YTpj56xP7y/aQ+5b/utxieQ25atcE55a9Z2rQBTj8GQIbHmcPO+OhXT17qXeo36S3vD7jutq1UpcM7zuTRzLDyrJXNTwXUxQRsH++SKBk5B8tVi7HeBWu7ntW3/he13qcEpJyWi8VXtknByTBHM3x/l1i4WOYXSOebI/FlFeDuVdw6pYqQbrvsPjLZQS0qEgxCIWHNHPajIzLJKBzLE/1EQqWQty8mgWx8OV4EVK8EMSIoVFijSNVUeQVUVQFGegAwNc2tVVVVVACooVQPJQAAoA+AAH2azr+kE70u80d7iobVuM1WOCjFNP4SV38WSaSXAc2IJyDGkKkBOTpJk5z0iAZcySr/wEWjRbUCQaxBv987iwI7HiG+AqsxxHTHQAk+VIe700/e9v3kuK6G4UY4N0mr159l26yoigqMJZnWQW3dp6szeIJAMqrKiqyYUFiTfdpkgcG8Q3+aK//eBo0H2XdBufJ7D7DZawXxhVETHnU+ayKQGQrhg4BHXWCt9m9lkL45zA/N+68P635T1+Hl6HTJ7QO2/edxUJue6W7UasHWORkSIMCCrGCvHDExUgMpeNirAMCCAR1/ZjwM24bnt+3onOb1yCBl98RbmssfhHAs0h9cIffqcw8c4zSXG/HyG3+2isP9Ijv7Nf2Co2c0+HasrA+klqRkdSPMMo29Cc/jasB9Fz2b132S/as14JWs7rKsLSRpIRDDVqRcoLocATe0ZC9M/HOqsd/wB3lZeLNxC5/QkNKkfd97rLOMfxvr19feNLngrvB79Qritt27W61YPJIIYhAVDO3NIw8WCVsu2WP1sZY68OgdLjNjaasA+5+6otul4EoDypVB/gIh/7NKLvaX46+xyxwokftU9eHCKEBAkE0nRQAeZIih+Day5n73PFxBCcQ7gXIwgIrYLHooIFUdCca0K76u4P7HtETnLuXmcn1ZIERiR85evz+WrGn6e5udAx7gbde1+G/itW7UZHwNKyH+PCW/1d37pe9znhXxt5WUjKUoZJc+55FMEf+iZSPivw1Gnv3MeAjBtrWnXEl9/E6+fgoWSAfI5eUfCTOo1a7RZYyc15vZvdH05+trA7HaccPTGAjvP75+vL+2lYWR9eJ4c7XKdm3JVruzvGhYuFPhtykBgjn8LBYdcYPoT114jvK9ophrrViYiWyGLsPNIhgN19DKWCD9qH8vPXlO6tw9h7NtsBEUQIT0A/uknXoAFHJn7dcwyNYcdQjwYBfi89B0Hn15rscGjtGnSZ85rwYOpvmfLpy5LqfpEO3n7nbKacDlb27+LVjKNyvDX5cXLKkEFSFZYEZfrLJOhGOViuSiIAAAMAdAB5ADpgfD3e4ab/AHr+2k7xvty2pJqxN7FQGens8LuolVckD2mTnsZABaN4g3VAB+u7n3Xtw36WytN4Yoaip41ifm5PEfqkCCMMWkKgyN+CEXkznxBjr2JG3Gg4n7eJK1bySdaMeoHu/wDzUqoHkP6eertf1qLef2V2z/J2Pzco/PqpHaPwRJQ3C7t8siSS0bDV5JYwQkjKqsWQN1A+tykEkhgRnpq/DkRTWGG6VFYXsTs+z8AcZWSCPbLlTak6dSsqU4GIz5gfdCUnH4h9RryncM4c8fi7Z/dWNy23ySjYjX5DxJo/yY9de04sQVuzPa4nGH3TfJZwR+nEc9qxFn5JWj/gDXd/RW8PLJv1+xj61TaxGPlZsrn7f0J+Q/HWG94EMzupI9APdFqhqtP0i3EHg8JbiP1zJSqn5S3IFb/RDfZqyo1Rz6WLiMptG11lPW1unO498cNSwxz8pXgPzA1r2GzjnYPP23RZpbZtM0ziOvBPNKQSIoIpJpGA/CKxQo7sFzkkKQB1ONd0ezDdv2H3n/Nt7/ldW6+ie2APu28WCAfZtvqQq3qDYnndsH9zUAOtOwNTuVqBgk4A29gefUWiyG7v3Zjdg4f47uWqVqup4ekqVzYglgZyUtS2QiTpG5UBa+WC8pPQElWA8h3HVQ8W7JzY6T2Sn7r2C2vT97zfZnWtXb1wgbmx7zUXPNb2u/An7qSrKiY+TFT89Ykdj3H5o7ltm5ICfY7leywAJYw8wjtIqjrzPXkmReh+sQMEkarjTHJZKRzPh+mkW+mOmsmO+/2E8QSb/ve5fcyzJtoMEkdtWgMSV4qFZHYr44lAjkSYsPCJzkjIIOtU+GOJ4LVeGzWlSWvYjSWGVCGSRGUMrqw6EEEeWceXn00mO/dxMK3Ce9OSQZqy0kI8w1uaKkGHUdV8fn6eik+moPDldDKKG523RYy1tiay8VVDhrk0NND54axKldD5jODID5jy9NaR/SpdmYO3bXuMagfc+dqUpz+DXsqgjwM4JWeCCMdCcSsemNUt7qGyCxxPsELLlX3BZSPhBBPdDD9ya4PyGtde852bfdLYN1orjxJqkjQMRnlnjxNXfHwljQ/Hy9dTWbP8OeM9B6HZFibwlwbbuWEq0a8li1IG8OCPHO/KMsQXZVVVGCWdlUZGSMjWmPcW7lU+0udz3UKNyliMMFZXDrQiblaTnkQlJbcvKqsULRwqCqM/NI8meHd77S/udvO07iWKRV7UZsZ+qBXmQ1rPig9AscczSHI+qYwfNRrcze92EdeaYkBYoZJeb0AWMvzfLA141OZ7ajHIosL+3XiT2nfd8sk9JN33EA/tIrUlaLPyigj/ACa67/sv3XAI2jdyCAwZdvuMpB8iGSuQQfgdcHhf9FXqYdOt/caiSL6k2rsSyA49SZmz8db/ANeAKAo8lAUfIAAayMrK/ZQxoF7eyLC3s47Gd2n3LbofuVuiiTcKSyPJRtRIkXtURmaSSaBERViEjEsw6DA6kA6/9tvYOm6vRMkpRKsshkUDJmicLzxKwZfDLGNBzYbC5wAcFWr4fz19NQ8moyOkbIzuubdV5rDy8SLLiMM7bYasdaNj1XGo0VRFRFCoihUUeSqBhQB6AADU65GjUXfVZYFChyVHu2viQz7lbfP1I28GMe5Y1Ct+V+c9PfrsO9Xx79xOD0qRMyX91U1EKkq6GVTJdnVh1Qww8yI36V2iXPVdeW2iMTXYufqJri82fUPP9bPwwdI76RTia9c4jkritZ9m2uCKvXxDKVd5Y47FiZHCFWVy0UfTOPZ/MEsBoXYfHGZnT5j+d+7r/C6b2rJx8fGwm8gLP6Rw/lVMdwqk9cKD0Ayce4AeZ9wAPUeWtpu5T2Kfcnh+pDInLbtD2+9nHMJ51VvCJwufZ4xHXHQdIs9SSTnL3LO7tY3LfaxsVp0obfIl228kUiJI0bBq1ZXkRQ7SzcjuE5wYopAeXnQnY2Py/p/J8Ndc1Wa6jHzK5miaQAEnyAJP2ddYC9o/EZsbhuVtm5vab9+yG96vZmkjP8Ar+TW5fbRvr19n3WxGGMlfbbs0YUFmLpWldAqrksxYDAAJJ6DWE3CPAk1l61OGCZzO8FRVRHb+2skHUhcgAPl2OOQDmOmkhreOR3gis53wImq8P8B7W3Ro9okuWF903g0kz6fprFoZIz+U6SXY925bns800+1zJFLYjSGXxIllSRFZnQMjEHKMzFSGGOZvPVhfpOqkq7zt0ZRvBh2iKOBuU4YieQTBTjB5eWLmA6jK5AyNVCWm58kc/JWP82pPGDHwji3uyfqSisj/AFxvi39e1P4nHpads3eM3befZvupPHIKni+AEiSIKZOQSFuT8IkRoAfQZ9+l+uzznygnP+Cf+ZTr7Jw3ZPlVsn5QSn/49X2wwtNgNB6ovXdk/bxuu0mwdrteB7SIxP8Ae438Tw+fwziVHwV8R/IjPMdMB+/rxYf77MP8Xrf/AEHSS/qWt/rK8flVnP5o9T/Upd/WG4fZTtH80J0fFE424NKK3XdW74HElziPa6dy+1ipblsQ2IWhgUcop2ZhIrRQo6sjQLj6xVg5UqSVIr13juy+Ta983Kk64RbMtiqcYD1J3eWsyAfpEVnrE4A568g6gadH0cnZzbbiiKaalcjgqULkxlmrTwoJGMEEKh5okUyMJpSqhuYorHGFOrzd7fui1t/rxESivuNVXFS1yKwYMDmtZGOd6zNiTEbI6OvMrdXVox88eNkU2g0jeuqLJjgntg3egCu3bpfrITnw4piYcnqWFaXxIAzHzZYuY+pOudx12+75uEPs+5btbs1+dJPBl8ERl0PNG7LDFDzFCAy8xYA9eXIBHqO0Huc8T0WIm2i1Ogx99oI12NjnH1VgU2AB5kyVowB6jqB4P/sm3jOPuJvf+a7+f9k1KtMJ7w4SfoifX0a3C/jcVQysuVo0btgH8WV1iqIfhmKzYH5fec66yDWfX0XXZDfrT7vevU7VZJYqdOstqCWvJIVexLYdIrEcb+GA9dOflwzBgM+Hk6D61TUZOKY1vsEWFXeY7NPufvu70Sp8FbUskQI6GCwPaYwo/EjWYQj0+9n460Nqdtb2uzO3e582k2O7QmbyItRRybe79ST1kAlXmOSpUnz0nPpUOyR1u0N3hidksQGhbZFLBJImaao8nKCEDJJYj5mIyURfdpT9lW7X24H4opxV7Lom4bXMuI3I5LE0EdtEVVJIRYBNJgEAS5Pu1LyFs8UTrFgj8H1RVx2fcnhlgmhYpLWmhsQOMZSWGVJoXAIIJR40bDAg4wQRkaeb9/Liw/32cfKCt/PBpLDha3+s7nv6Vpj+aPQeFLn6xv8A2VLJ/NCdSj2RyG3gH+SJzHv18WfsxL/kKv8Ay2tTO6vxdbucO7NcvSeJatUIZ5pOUJ4hccyyFEwql0KsQoAyemsTH4N3BgRHtu5M7dEUUbf1mPRV6wAdTgdSB8Rrdzsn4VFTbNuqBeUVKNWsF/FEUCR4+zlxqC1RsbGNDAAb8EXrBqdQNTqARZ83w0Vh+U4eGduU+oaOY8v5OQH5E+/Vp+FO8Nt0kCtPIsUwH3xGVvP1KMFIcHzGDkZwQMaTXeE4GaveeYD7xaPOrDyWTCiRD8Sfrr7wT+L1Vx1wuLPytEyJYmgc+R5EDkQu8y4GLruNFM8m6G7eYNbg3f8A1XDp94jaWkWMTsCxChjFIqZPQZcoAB5dTgaZ0bDHT166oJwpwtJasR14lJMjAMfREz9d29wVQf32F89X2qRYVV/FAH5BjXROz2qZOoNe+ZoAHIgVfl9FzXtFpWPpz2RwvJJBsGtq5ch4r7soIwR018YdvjXqqIPkoH5hr7jUO+PlrcFqC481WNvwlRsdOoBx5dOoOPTXzbboR/c4x+9X/h8DqqHcs4rik3TiDw7Mcy7o8W/DkkR/Caa3uFDwiiktHivSpNh8H62MdDpudvR/RnDX/rqf7vv/ANPt0RNaOrH6Knr5BfTofyeR92v2OQdPqg/YNKPuxA+x384/7+33/edj3H7ft0su1bu+bRY442OzYph559v3S9K5klHPY2+bY49um5VkChqqzzBQFCnxCWDHRFa4gD3ahJlPkQfTp1+Y6eulh3ot9mg4e3eSvKYp/YpI4px5wNLiETjqOsPieIOo6rpWd3vs0qbNxHuO07XG0W3S7DtO6NAZJJFFxrW4UZbGZmc89qGrB4mG+s0JYgFiSRWidwPPH2/06akEHSB70PB0W4WOH9rtknb7u5TSX4Azp7ZHWoWZ4qsjxujeEZjFM65w3ghT56c3BvC0VSpXqQmQw1YY68XiO0j8kahE55HJaRgoALOSzYySToi7gOM4yM6l3A8+mqecF8Xhu0G7KLEXLPFa2I1xKhfNKptm5VpzEMugElreIDzcoJSNsHnUhr99oH+pLiHAyfuVb6dfrfez9XoQevl0IOqInUrA9Rj+n9D/AC/HQZ194z89Vy7jdKJdsuGrDNX21t43BdtoT83jbdDE61pq0qvJM0ZNqK1OkbPlI51UBQBrqe3DgunZ4t2Q29nsXzBRaWvNE6qu1yDca7C7OGswFkBRQORZm6EcmCdKRWilhUj6ygj4gH+Q6/FaGPH1AnL+1Ax7vTp8NdD2kbXZm22/DTlEdualaiqy/qU7QukMnp+A5VvTy0m+4vYojZmgqbfNt9ipaevuu3zM7vV3BYYPaCrOz88NkeHajkQ8kqzeJgF21VFYgRj3DU41OjRFGNAGp0aIjRo0aIuu3rZoZo2jmjSSNhhkcAg/YdLSbu37SWJ8GUdfwRNLj/XJ+zONTo1C52JBMWmVjXfMAqRxcyeAERSOb8iQvZcHcAU6gK1YFTmALNlmdvcGd2ZiB6DOBr04QajRqWjiZE0NYAB0AoLBfI+Rxe8knqTZ9V+9dHxtMy07bIxV1qzlWHmpETkMMgjIIBGQdTo1cXhV/wCx/sS27bdz2A7fXWAtw5brT+GqL7UEk2l45bRVFM9hCXxK55vv0mc+Ideq75O3Mdis2YZ5YLW2vDfp2IvD54Z43wDiaKaN0dXeN1eNgyOR089To0Rd93Z+Eo6uyUUR5HaZJbs8spUyT2Lc0ly1LIUSNcyTWJGwiIqjCgAKNfvi3YUbiHZrBLeJDt29wqARylZZdodiwwTzAwJy4YDBbIbIxOjRF6ftP4Kgvbdeo2lY17lWevMFPK3JJGyMVbryuM5U4OCAeukF3Gdoklr39yuW7VvcJ7B21rE5iBFWhJOtSFErwV4xymxPI78heR5SWY4ULOjRF33fW2InaFuwzzQXtospeoWYfD8SGUo9VwRNFNG8ckNmVHR42DZB81B02+zWqybfRR5pZnFSvzzTFWlmYopeSVkWNS7klm5ERcnoqjpqNGiKvfDnY9Rjba9yWJfuk/El6xJcKRCeb2iXdY5IJZUiVnrqjKiJ+lEaZLFSS3e8/sCWOH94ryFwk9CeJyhAYK68rcpZWAOCcEqce7UaNEX47ItqEV/iJFdyjbtHOqMQRG0210JZRHgAhHk5pMEseZ264ONLfvCtZj3/AGeares1y0KQTpEK7JZha9C7QzCxWnYKxGMwtC4DEBhnRo0RP/j2mz0LiJNLC71LCpNCVEsLGJwssLOkiiWM4dC6OoYDKsMgp7uTbe52c3rFiexd3WzJcu2JvDDSSKkdRAqQRQRRRpFWiVUjiUZBY5LE6NGiKwGjRo0RGjRo0RGjRo0Rf//Z"
PASSWORD_ADMIN = "alegre2026"
SERVICIOS_GRID = ["Pilotos","Pescadores","Logista-Saludes","Colegio","Transfer",
                  "Boda Ida","Boda Regreso 1","Boda Regreso 2"]
PRECIOS_FIJOS = {"Pilotos":15,"Pescadores":15,"Logista-Saludes":15,"Colegio":25,
                 "Transfer":20,"Boda Ida":40,"Boda Regreso 1":40,"Boda Regreso 2":40}
PRECIO_DIETA=15; PRECIO_HORA_SEMANA=10; PRECIO_HORA_FINDE=12
FESTIVOS = {
    date(2025,1,1),date(2025,1,6),date(2025,3,19),date(2025,4,17),date(2025,4,18),
    date(2025,4,28),date(2025,5,1),date(2025,8,15),date(2025,10,9),date(2025,10,12),
    date(2025,11,1),date(2025,12,6),date(2025,12,8),date(2025,12,25),
    date(2026,1,1),date(2026,1,6),date(2026,3,19),date(2026,4,2),date(2026,4,3),
    date(2026,4,20),date(2026,5,1),date(2026,8,15),date(2026,10,9),date(2026,10,12),
    date(2026,11,1),date(2026,12,6),date(2026,12,8),date(2026,12,25),
}
COLUMNAS_BD = ["conductor","fecha","dieta","servicios_fijos","servicios_horas","observaciones","hoja_servicio","tipo_registro"]

CSS = """
<style>
header{visibility:hidden;}
#MainMenu{visibility:hidden;}
footer{visibility:hidden;}
[data-testid="stToolbar"]{display:none;}
[data-testid="stDecoration"]{display:none;}
[data-testid="stSidebar"]{display:none !important;}
[data-testid="manage-app-button"]{display:none !important;}
.stAppDeployButton{display:none !important;}
.block-container{
    padding-top:0.4rem !important;
    padding-bottom:0.4rem !important;
    max-width:460px !important;
}
.stButton>button{
    height:2.8rem;font-size:0.95rem;
    font-weight:700;border-radius:10px;width:100%;
}
.stSelectbox>label,.stDateInput>label,.stTextInput>label,
.stTimeInput>label,.stNumberInput>label,
.stCheckbox>label span,.stTextArea>label{
    font-size:0.9rem !important;font-weight:600 !important;
}
div[data-testid="stVerticalBlock"]>div{margin-bottom:-0.3rem;}
.admin-link{text-align:center;font-size:0.65rem;margin-top:1rem;opacity:0.25;}
</style>
"""
CSS_ADMIN = """
<style>
[data-testid="stSidebar"]{display:block !important;}
.block-container{max-width:1200px !important;padding-top:1rem !important;}
</style>
"""

def logo_html(w="160px"):
    return f'''<div style="text-align:center;margin-bottom:0.2rem;">
<img src="data:image/jpeg;base64,{LOGO_B64}" style="max-width:{w};width:50%;border-radius:6px;">
</div>'''

def limpiar_pdf(texto):
    """Elimina caracteres no compatibles con PDF Latin-1."""
    if not texto: return ""
    reemplazos = {
        "a":"a","e":"e","i":"i","o":"o","u":"u",
        "A":"A","E":"E","I":"I","O":"O","U":"U",
        "n":"n","N":"N","u":"u","U":"U",
        "EUR":"EUR","-":"-","->":"->",
        "(!!)":"(!!)","OK":"OK","X":"X",
        " ":"",
    }
    tabla = str.maketrans({
        "\u00e1":"a","\u00e9":"e","\u00ed":"i","\u00f3":"o","\u00fa":"u",
        "\u00c1":"A","\u00c9":"E","\u00cd":"I","\u00d3":"O","\u00da":"U",
        "\u00f1":"n","\u00d1":"N","\u00fc":"u","\u00dc":"U",
        "\u20ac":"EUR","\u2014":"-","\u2192":"->","\u2014":"-",
    })
    texto = texto.translate(tabla)
    texto = texto.encode("latin-1","replace").decode("latin-1")
    return texto

def fijos_a_texto(raw):
    try:
        items = json.loads(str(raw or "[]"))
        partes = [f"{i['tipo']} x{i['cantidad']}" if int(i.get("cantidad",1))>1 else i["tipo"]
                  for i in items if i.get("tipo")]
        return ", ".join(partes)
    except: return str(raw) if raw else ""

def horas_a_texto(raw):
    try:
        items = json.loads(str(raw or "[]"))
        partes = [f"{i['concepto']} {i.get('inicio','')} a {i.get('fin','')}"
                  if i.get("concepto") else "" for i in items]
        return ", ".join(p for p in partes if p)
    except: return str(raw) if raw else ""

def necesita_revision(row):
    sf    = fijos_a_texto(row.get("servicios_fijos",""))
    sh_raw = row.get("servicios_horas","")
    sh    = horas_a_texto(sh_raw)
    if not sf and not sh: return "Sin servicios registrados"
    try: json.loads(str(row.get("servicios_fijos") or "[]"))
    except: return "JSON fijos incorrecto"
    try:
        items = json.loads(str(sh_raw or "[]"))
        for item in items:
            concepto = item.get("concepto","").strip()
            inicio   = item.get("inicio","").strip()
            fin      = item.get("fin","").strip()
            # Si hay concepto pero falta el horario → revisar
            if concepto and (not inicio or not fin or inicio in ("","00:00") or fin in ("","00:00")):
                return f"Horario incompleto en: {concepto}"
    except: return "JSON horas incorrecto"
    return ""

def periodo_actual():
    """
    Calcula el periodo de revisión actual por defecto.
    Los periodos van del 16 de un mes al 15 del siguiente.
    """
    hoy = date.today()
    if hoy.day >= 16:
        # Del 16 de este mes al 15 del mes siguiente
        fecha_ini = date(hoy.year, hoy.month, 16)
        if hoy.month == 12:
            fecha_fin = date(hoy.year + 1, 1, 15)
        else:
            fecha_fin = date(hoy.year, hoy.month + 1, 15)
    else:
        # Del 16 del mes pasado al 15 de este mes
        if hoy.month == 1:
            fecha_ini = date(hoy.year - 1, 12, 16)
        else:
            fecha_ini = date(hoy.year, hoy.month - 1, 16)
        fecha_fin = date(hoy.year, hoy.month, 15)
    return fecha_ini, fecha_fin

@st.cache_resource
def get_supabase():
    return create_client(st.secrets["supabase"]["url"], st.secrets["supabase"]["key"])

@st.cache_data(ttl=60)
def cargar_conductores():
    """
    Devuelve lista de strings para el desplegable.
    Si el conductor tiene DNI: "García López, Antonio — 12345678A"
    Si no tiene DNI: "García López, Antonio"
    Si la columna dni no existe aún, carga solo nombres sin error.
    """
    try:
        # Intentar cargar con DNI
        r = get_supabase().table("conductores").select("nombre,dni").order("nombre").execute()
        lista = []
        for f in r.data:
            if not f.get("nombre"): continue
            nombre = f["nombre"].strip()
            dni    = (f.get("dni") or "").strip()
            lista.append(f"{nombre} — {dni}" if dni else nombre)
        return lista
    except Exception:
        try:
            # Si falla (columna dni no existe), cargar solo nombres
            r = get_supabase().table("conductores").select("nombre").order("nombre").execute()
            return [f["nombre"].strip() for f in r.data if f.get("nombre")]
        except Exception as e:
            st.error(f"Error al cargar conductores: {e}")
            return []

def extraer_nombre(seleccion):
    """Extrae solo el nombre sin el DNI del string del desplegable."""
    if " — " in seleccion:
        return seleccion.split(" — ")[0].strip()
    return seleccion.strip()

@st.cache_data(ttl=20)
def cargar_datos_brutos():
    try:
        # Incluimos el ID para poder eliminar registros específicos
        cols_con_id = "id," + ",".join(COLUMNAS_BD)
        r = get_supabase().table("datos_brutos").select(cols_con_id).order("id").execute()
        if not r.data: return pd.DataFrame(columns=["id"]+COLUMNAS_BD)
        df = pd.DataFrame(r.data)
        for col in COLUMNAS_BD:
            if col not in df.columns: df[col] = None
        return df[["id"]+COLUMNAS_BD]
    except Exception as e: st.error(f"Error: {e}"); return pd.DataFrame(columns=["id"]+COLUMNAS_BD)

def eliminar_por_ids(ids_lista):
    """Elimina registros específicos de Supabase por su ID."""
    try:
        sb = get_supabase()
        for i in range(0, len(ids_lista), 500):
            sb.table("datos_brutos").delete().in_("id", ids_lista[i:i+500]).execute()
        cargar_datos_brutos.clear()
        return True
    except Exception as e:
        st.error(f"Error al eliminar: {e}"); return False

def guardar_fila(d):
    try:
        get_supabase().table("datos_brutos").insert(d).execute()
        cargar_datos_brutos.clear(); return True
    except Exception as e: st.error(f"Error: {e}"); return False

def guardar_tabla_completa(df):
    try:
        sb = get_supabase()
        ids = sb.table("datos_brutos").select("id").execute()
        if ids.data:
            all_ids = [f["id"] for f in ids.data]
            for i in range(0,len(all_ids),500):
                sb.table("datos_brutos").delete().in_("id",all_ids[i:i+500]).execute()
        if not df.empty:
            filas = []
            for _,row in df.iterrows():
                f = {col:(str(row[col]) if row[col] is not None else None)
                     for col in COLUMNAS_BD if col in df.columns}
                if "dieta" in f: f["dieta"] = str(f["dieta"]).lower() in ["true","1","si"]
                # No incluir el campo id al reinsertar
                f.pop("id", None)
                filas.append(f)
            for i in range(0,len(filas),500):
                sb.table("datos_brutos").insert(filas[i:i+500]).execute()
        cargar_datos_brutos.clear(); return True
    except Exception as e: st.error(f"Error: {e}"); return False

def es_dia_especial(fecha):
    if isinstance(fecha,datetime): fecha=fecha.date()
    elif isinstance(fecha,str):
        try: fecha=pd.to_datetime(fecha).date()
        except: return False
    return fecha.weekday()>=5 or fecha in FESTIVOS

def calcular_horas(h1,h2):
    try:
        if isinstance(h1,str): h1=datetime.strptime(h1.strip(),"%H:%M").time()
        if isinstance(h2,str): h2=datetime.strptime(h2.strip(),"%H:%M").time()
        i=datetime.combine(date.today(),h1); f=datetime.combine(date.today(),h2)
        return round((f-i).total_seconds()/3600,4) if f>i else 0.0
    except: return 0.0

def calcular_liquidacion(df):
    if df.empty: return pd.DataFrame()
    df = df.copy()
    # Solo calcular registros de tipo "extra" (no jornada habitual)
    if "tipo_registro" in df.columns:
        df = df[df["tipo_registro"].fillna("extra") != "jornada"]
    if df.empty: return pd.DataFrame()
    df["fecha"] = pd.to_datetime(df["fecha"],errors="coerce").dt.date
    def a_bool(v):
        if isinstance(v,bool): return v
        return str(v).strip().lower() in ["true","1","si"]
    df["dieta"] = df["dieta"].fillna(False).apply(a_bool)
    resultados = []
    for _,fila in df.iterrows():
        conductor=fila.get("conductor",""); fecha=fila.get("fecha")
        if pd.isna(fecha): continue
        tarifa = PRECIO_HORA_FINDE if es_dia_especial(fecha) else PRECIO_HORA_SEMANA
        imp_dieta = PRECIO_DIETA if fila.get("dieta") else 0.0
        fijos_txt=[]; horas_txt=[]; horas_tot=0.0; imp_horas=0.0; imp_fijos=0.0
        sf_raw=fila.get("servicios_fijos")
        if sf_raw and str(sf_raw) not in ("","nan","None","[]"):
            try:
                for item in json.loads(str(sf_raw)):
                    tipo=item.get("tipo",""); cant=int(item.get("cantidad",1))
                    if tipo in PRECIOS_FIJOS:
                        imp_fijos+=PRECIOS_FIJOS[tipo]*cant
                        fijos_txt.append(f"{tipo}" if cant==1 else f"{tipo} x{cant}")
            except: pass
        sh_raw=fila.get("servicios_horas")
        if sh_raw and str(sh_raw) not in ("","nan","None","[]"):
            try:
                for item in json.loads(str(sh_raw)):
                    c=item.get("concepto",""); h1=item.get("inicio",""); h2=item.get("fin","")
                    if h1 and h2 and h1!="00:00":
                        h=calcular_horas(h1,h2); horas_tot+=h; imp_horas+=h*tarifa
                    if c: horas_txt.append(f"{c} {h1} a {h2}" if (h1 and h2 and h1!="00:00") else c)
            except: pass
        obs=str(fila.get("observaciones","")) if fila.get("observaciones") else ""
        revision=necesita_revision(fila)
        nro_s=fila.get("hoja_servicio","")
        resultados.append({
            "Nº Serv.":int(nro_s) if nro_s and str(nro_s) not in ("","nan","None") else "",
            "Conductor":conductor,"Fecha":fecha,
            "Servicios Fijos":" | ".join(fijos_txt) if fijos_txt else "—",
            "Por Horas":" | ".join(horas_txt) if horas_txt else "—",
            "Observaciones":obs,
            "Horas Tot.":round(horas_tot,2),
            "IMP. HORAS":round(imp_horas,2),
            "IMP. FIJOS":round(imp_fijos,2),
            "IMP. DIETAS":round(imp_dieta,2),
            "TOTAL":round(imp_horas+imp_fijos+imp_dieta,2),
            "REVISAR":"REVISAR: "+revision if revision else "",
        })
    return pd.DataFrame(resultados).sort_values(["Conductor","Fecha"]).reset_index(drop=True)

def generar_excel_bytes(df_cierre):
    output=io.BytesIO()
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Liquidacion"
    C_AZUL="1F4E79"; C_AMAR="FFF2CC"; C_VERDE="C6EFCE"; C_FIJO="EBF5FB"
    C_HORA="FEF9E7"; C_REV="FADBD8"
    borde=Border(
        left=Side(style="thin",color="CCCCCC"),right=Side(style="thin",color="CCCCCC"),
        top=Side(style="thin",color="CCCCCC"),bottom=Side(style="thin",color="CCCCCC"))
    al_c=Alignment(horizontal="center",vertical="center",wrap_text=True)
    al_d=Alignment(horizontal="right",vertical="center")
    al_i=Alignment(horizontal="left",vertical="center",wrap_text=True)
    cabs=["Nº Serv.","Conductor","Fecha","Servicios / Horario","Obs./Revisar",
          "Horas","IMP. HORAS","IMP. FIJOS","IMP. DIETAS","TOTAL DIA"]
    for cn,t in enumerate(cabs,1):
        c=ws.cell(row=1,column=cn,value=t)
        c.font=Font(bold=True,color="FFFFFF",size=10,name="Calibri")
        c.fill=PatternFill("solid",fgColor=C_AZUL); c.alignment=al_c; c.border=borde
    ws.row_dimensions[1].height=28
    fila_xls=2; conductor_actual=None
    sub_fijos=0; sub_horas=0; sub_dietas=0; sub_total=0
    def subtotal(fxls,nombre,sf,sh,sd,st_):
        c=ws.cell(row=fxls,column=1,value=f"SUBTOTAL — {nombre}")
        c.font=Font(bold=True,size=10,color="FFFFFF",name="Calibri")
        c.fill=PatternFill("solid",fgColor="1A5276"); c.alignment=al_i; c.border=borde
        for cn in(2,3,4,5,6):
            cx=ws.cell(row=fxls,column=cn)
            cx.fill=PatternFill("solid",fgColor="1A5276"); cx.border=borde
        for cn,val in zip((7,8,9,10),(sh,sf,sd,st_)):
            cx=ws.cell(row=fxls,column=cn,value=round(val,2))
            cx.number_format='#,##0.00 "€"'; cx.font=Font(bold=True,size=10,color="FFFFFF",name="Calibri")
            cx.fill=PatternFill("solid",fgColor="1A5276"); cx.alignment=al_d; cx.border=borde
        ws.row_dimensions[fxls].height=20
    for _,fila in df_cierre.iterrows():
        conductor=fila["Conductor"]; fecha=fila["Fecha"]
        fondo=es_dia_especial(fecha)
        fill_f=PatternFill("solid",fgColor=C_AMAR if fondo else C_FIJO)
        fill_h=PatternFill("solid",fgColor=C_AMAR if fondo else C_HORA)
        revisar=fila.get("REVISAR") or ""
        if not revisar or str(revisar).strip() in ("None","nan"): revisar=""
        fill_rev=PatternFill("solid",fgColor=C_REV) if revisar else None
        if conductor_actual is not None and conductor!=conductor_actual:
            subtotal(fila_xls,conductor_actual,sub_fijos,sub_horas,sub_dietas,sub_total)
            fila_xls+=1; sub_fijos=0; sub_horas=0; sub_dietas=0; sub_total=0
        conductor_actual=conductor
        sub_fijos+=fila["IMP. FIJOS"]; sub_horas+=fila["IMP. HORAS"]
        sub_dietas+=fila["IMP. DIETAS"]; sub_total+=fila["TOTAL"]
        dieta_label="Dieta" if fila["IMP. DIETAS"]>0 else ""
        obs_col=revisar if revisar else dieta_label
        nro_s=fila.get("Nº Serv.","") or ""
        vals1=[nro_s,conductor,fecha,fila["Servicios Fijos"],obs_col,
               "",None,fila["IMP. FIJOS"],fila["IMP. DIETAS"],fila["TOTAL"]]
        for cn,val in enumerate(vals1,1):
            c=ws.cell(row=fila_xls,column=cn,value=val); c.border=borde
            c.fill=fill_rev if (revisar and cn in(4,5)) else fill_f
            if cn==1: c.alignment=al_c; c.font=Font(size=10,name="Calibri")
            elif cn==2: c.font=Font(bold=True,size=10,name="Calibri"); c.alignment=al_i
            elif cn==3: c.number_format="DD/MM/YYYY"; c.alignment=al_c; c.font=Font(size=10,name="Calibri")
            elif cn in(4,5):
                c.alignment=al_i
                c.font=Font(size=9,name="Calibri",bold=bool(revisar and cn==5),
                            color="C0392B" if (revisar and cn==5) else "000000")
            elif cn in(8,9,10):
                c.number_format='#,##0.00 "€"'; c.alignment=al_d
                c.font=Font(bold=(cn==10),size=10,name="Calibri")
                if cn==10 and not (revisar or fondo): c.fill=PatternFill("solid",fgColor=C_VERDE)
        ws.row_dimensions[fila_xls].height=18; fila_xls+=1
        por_horas=fila["Por Horas"]; obs=fila["Observaciones"]
        if por_horas!="—" or obs:
            contenido_h=por_horas if por_horas!="—" else ""
            vals2=["","","",contenido_h,obs,
                   fila["Horas Tot."] if fila["Horas Tot."]>0 else "",
                   fila["IMP. HORAS"] if fila["IMP. HORAS"]>0 else None,"","",""]
            for cn,val in enumerate(vals2,1):
                c=ws.cell(row=fila_xls,column=cn,value=val); c.border=borde; c.fill=fill_h
                if cn in(4,5): c.alignment=al_i; c.font=Font(size=9,italic=True,name="Calibri")
                elif cn==6: c.alignment=al_c; c.font=Font(size=9,name="Calibri")
                elif cn==7: c.number_format='#,##0.00 "€"'; c.alignment=al_d; c.font=Font(size=9,italic=True,name="Calibri")
            ws.row_dimensions[fila_xls].height=16; fila_xls+=1
    if conductor_actual:
        subtotal(fila_xls,conductor_actual,sub_fijos,sub_horas,sub_dietas,sub_total)
    anchos={1:8,2:24,3:12,4:42,5:28,6:8,7:14,8:14,9:13,10:14}
    for cn,ancho in anchos.items():
        ws.column_dimensions[get_column_letter(cn)].width=ancho
    ws.freeze_panes="A2"; ws.auto_filter.ref="A1:J1"
    wb.save(output); output.seek(0); return output.getvalue()

def generar_pdf_bytes(df_cierre):
    pdf=FPDF(orientation="L",unit="mm",format="A4")
    pdf.set_auto_page_break(auto=True,margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica","B",16)
    pdf.set_text_color(31,78,121)
    pdf.cell(0,10,limpiar_pdf("AUTOCARES ALEGRE"),new_x="LMARGIN",new_y="NEXT",align="C")
    pdf.set_font("Helvetica","",11)
    pdf.set_text_color(80,80,80)
    pdf.cell(0,6,limpiar_pdf(f"Resumen de Liquidacion  -  {date.today().strftime('%d/%m/%Y')}"),
             new_x="LMARGIN",new_y="NEXT",align="C")
    pdf.ln(3)
    total_g=df_cierre["TOTAL"].sum()
    pdf.set_font("Helvetica","B",10)
    pdf.set_fill_color(235,245,251); pdf.set_text_color(0,0,0)
    pdf.cell(0,7,limpiar_pdf(f"  TOTAL GENERAL: {total_g:,.2f} EUR   |   {df_cierre['Conductor'].nunique()} conductores"),
             fill=True,new_x="LMARGIN",new_y="NEXT")
    pdf.ln(4)
    for conductor,grupo in df_cierre.groupby("Conductor"):
        pdf.set_font("Helvetica","B",11)
        pdf.set_fill_color(31,78,121); pdf.set_text_color(255,255,255)
        pdf.cell(0,8,limpiar_pdf(f"  {conductor}"),fill=True,new_x="LMARGIN",new_y="NEXT")
        pdf.set_font("Helvetica","B",8)
        pdf.set_fill_color(215,230,242); pdf.set_text_color(0,0,0)
        pdf.cell(25,6,limpiar_pdf("Fecha"),fill=True,border=1)
        pdf.cell(95,6,limpiar_pdf("Servicios Fijos"),fill=True,border=1)
        pdf.cell(85,6,limpiar_pdf("Por Horas / Observaciones"),fill=True,border=1)
        pdf.cell(20,6,limpiar_pdf("Horas"),fill=True,border=1,align="C")
        pdf.cell(30,6,limpiar_pdf("Total"),fill=True,border=1,align="R",new_x="LMARGIN",new_y="NEXT")
        pdf.set_font("Helvetica","",8)
        for _,row in grupo.iterrows():
            fecha_str=row["Fecha"].strftime("%d/%m/%Y") if hasattr(row["Fecha"],"strftime") else str(row["Fecha"])
            fijos=limpiar_pdf(str(row["Servicios Fijos"])[:90]) if row["Servicios Fijos"]!="—" else ""
            obs_h = str(row["Observaciones"]) if row.get("Observaciones") else ""
            horas_txt = limpiar_pdf(str(row["Por Horas"])[:75]) if row["Por Horas"]!="—" else ""
            horas_col = (horas_txt + " | " + limpiar_pdf(obs_h) if horas_txt and obs_h
                        else horas_txt or limpiar_pdf(obs_h))[:85]
            revisar=row.get("REVISAR") or ""
            if not revisar or str(revisar).strip() in ("None","nan"): revisar=""
            bg=(252,215,215) if revisar else (255,255,255)
            pdf.set_fill_color(*bg)
            pdf.cell(25,5,fecha_str,fill=True,border=1)
            pdf.cell(95,5,fijos,fill=True,border=1)
            pdf.cell(85,5,horas_col,fill=True,border=1)
            pdf.cell(20,5,str(row["Horas Tot."]) if row["Horas Tot."]>0 else "",fill=True,border=1,align="C")
            pdf.cell(30,5,limpiar_pdf(f"{row['TOTAL']:,.2f} EUR"),fill=True,border=1,align="R",new_x="LMARGIN",new_y="NEXT")
            if revisar and str(revisar).strip() not in ("","None","nan"):
                pdf.set_font("Helvetica","I",7); pdf.set_text_color(192,57,43)
                pdf.cell(0,4,limpiar_pdf(f"  {revisar}"),new_x="LMARGIN",new_y="NEXT")
                pdf.set_font("Helvetica","",8); pdf.set_text_color(0,0,0)
            if row["Observaciones"]:
                pdf.set_font("Helvetica","I",7); pdf.set_text_color(100,100,100)
                pdf.cell(0,4,limpiar_pdf(f"  Obs: {str(row['Observaciones'])[:80]}"),new_x="LMARGIN",new_y="NEXT")
                pdf.set_font("Helvetica","",8); pdf.set_text_color(0,0,0)
        pdf.set_font("Helvetica","B",9)
        pdf.set_fill_color(196,230,197)
        sub=grupo["TOTAL"].sum()
        h_s=grupo["IMP. HORAS"].sum(); f_s=grupo["IMP. FIJOS"].sum(); d_s=grupo["IMP. DIETAS"].sum()
        pdf.cell(0,6,limpiar_pdf(f"  Subtotal {conductor}:  Fijos {f_s:,.2f}  +  Horas {h_s:,.2f}  +  Dietas {d_s:,.2f}  =  {sub:,.2f} EUR"),
                 fill=True,new_x="LMARGIN",new_y="NEXT")
        pdf.ln(4)
    buf=io.BytesIO(); pdf.output(buf); buf.seek(0); return buf.getvalue()


# ============================================================
# IMPORTACIÓN DESDE EXCEL
# ============================================================

SERVICIOS_VALIDOS = set(PRECIOS_FIJOS.keys())

def normalizar_servicio(texto):
    if not texto or str(texto).strip() in ("","nan","None"): return []
    mapa = {
        "logista":"Logista-Saludes","saludes":"Logista-Saludes",
        "logista-saludes":"Logista-Saludes","pilotos":"Pilotos",
        "pescadores":"Pescadores","colegio":"Colegio","transfer":"Transfer",
        "boda ida":"Boda Ida","boda regreso 1":"Boda Regreso 1",
        "boda regreso 2":"Boda Regreso 2",
    }
    return [mapa[p.strip().lower()] for p in str(texto).split(",")
            if mapa.get(p.strip().lower())]

def importar_conductores_excel(df_raw):
    sb = get_supabase()
    existentes = {c.split(" — ")[0].strip() for c in cargar_conductores()}
    importados = 0; omitidos = 0; errores = []
    for _, row in df_raw.iterrows():
        nombre = str(row.get("Nombre","")).strip()
        dni    = str(row.get("mm","")).strip()
        if not nombre or nombre == "nan": continue
        if nombre in existentes: omitidos += 1; continue
        try:
            dato = {"nombre": nombre}
            if dni and dni not in ("nan","None",""): dato["dni"] = dni
            sb.table("conductores").insert(dato).execute()
            importados += 1
        except Exception as e:
            errores.append(f"{nombre}: {str(e)[:60]}")
    cargar_conductores.clear()
    return importados, omitidos, errores

def importar_servicios_excel(df_raw):
    sb = get_supabase()
    importados = 0; errores = []
    for idx, row in df_raw.iterrows():
        try:
            nombre = str(row.iloc[2]).strip() if pd.notna(row.iloc[2]) else ""
            if not nombre or nombre == "nan": continue
            fecha_raw = row.iloc[4]
            if pd.isna(fecha_raw): continue
            fecha_str = (fecha_raw.date().strftime("%Y-%m-%d")
                        if hasattr(fecha_raw,"date") else str(fecha_raw)[:10])
            try: dieta = float(row.iloc[15] or 0) > 0
            except: dieta = False
            # Servicios fijos
            conteo = {}
            for ci in [9,10,11,12]:
                for s in normalizar_servicio(row.iloc[ci]):
                    conteo[s] = conteo.get(s,0) + 1
            fijos_ok = [{"tipo":t,"cantidad":c} for t,c in conteo.items()]
            # Horas → van a observaciones
            horas_txt = str(row.iloc[6]).strip() if pd.notna(row.iloc[6]) else ""
            horas_raw = str(row.iloc[7]).strip() if pd.notna(row.iloc[7]) else ""
            horas_info = horas_txt if horas_txt not in ("nan","") else (horas_raw if horas_raw not in ("nan","0","0.0","") else "")
            # Observaciones originales
            obs_orig = str(row.iloc[17]).strip() if pd.notna(row.iloc[17]) else ""
            if obs_orig in ("nan","None"): obs_orig = ""
            # Unir: horas por un lado, obs original por otro
            partes = []
            if horas_info: partes.append(f"[HORAS: {horas_info}]")
            if obs_orig:   partes.append(obs_orig)
            obs_final = " | ".join(partes) if partes else None
            sb.table("datos_brutos").insert({
                "conductor":       nombre,
                "fecha":           fecha_str,
                "dieta":           dieta,
                "servicios_fijos": json.dumps(fijos_ok, ensure_ascii=False),
                "servicios_horas": json.dumps([], ensure_ascii=False),
                "observaciones":   obs_final,
            }).execute()
            importados += 1
        except Exception as e:
            errores.append(f"Fila {idx}: {str(e)[:80]}")
    cargar_datos_brutos.clear()
    return importados, errores

def formulario_conductor(prefix=""):
    """Formulario reutilizable para conductor y admin manual."""
    conductores=cargar_conductores()
    busqueda=st.text_input("🔍 Escribe tu nombre",
                          placeholder="Escribe las primeras letras...",
                          key=f"{prefix}busqueda")
    cond_f=([c for c in conductores if busqueda.strip().lower() in c.lower()]
            if busqueda.strip() else conductores)
    nombre_sel=st.selectbox("👤 Selecciona tu Nombre",
                            options=["— Selecciona —"]+cond_f+["✏️ Mi nombre no aparece en la lista"],
                            key=f"{prefix}nombre")
    if nombre_sel == "✏️ Mi nombre no aparece en la lista":
        nombre_manual = st.text_input(
            "✏️ Escribe tu nombre completo",
            placeholder="Nombre y apellidos",
            key=f"{prefix}nombre_manual"
        )
        if nombre_manual.strip():
            nombre_sel = nombre_manual.strip()
        else:
            nombre_sel = "— Selecciona —"
    fecha_sel=st.date_input("📅 Fecha del Servicio",value=date.today(),
                            format="DD/MM/YYYY",key=f"{prefix}fecha")
    nro_servicio=st.number_input("📋 Nº Servicio",min_value=0,max_value=99999,
                                  value=0,step=1,key=f"{prefix}nro_servicio")
    st.markdown("---")

    # ── Tipo de registro ──────────────────────────────────
    tipo=st.radio(
        "¿Es un servicio extra (genera pago adicional)?",
        options=["➕  Sí — es un servicio EXTRA","📋  No — está dentro de mi jornada habitual"],
        key=f"{prefix}tipo_registro",
        horizontal=False
    )
    es_extra = tipo.startswith("➕")

    # Valores vacíos por defecto para jornada
    conteo={s:0 for s in SERVICIOS_GRID}
    horas_cap=[]
    dieta=False; dieta_conc=""; obs=""

    if es_extra:
        st.markdown("---")
        st.markdown("**📌 Servicios Fijos** — ¿Cuántos de cada uno?")
        st.caption("Pon 0 si no has hecho ese servicio hoy.")
        c1,c2=st.columns(2)
        for idx,s in enumerate(SERVICIOS_GRID):
            col=c1 if idx%2==0 else c2
            conteo[s]=col.number_input(s,min_value=0,max_value=5,value=0,
                                       key=f"{prefix}fijo_{idx}",step=1)
        st.markdown("---")
        st.markdown("**🕐 Servicios Por Horas**")
        st.caption("Si el horario no es exacto, escríbelo en Observaciones.")
        n_key=f"{prefix}n_horas"
        if n_key not in st.session_state: st.session_state[n_key]=1
        for i in range(st.session_state[n_key]):
            if i>0: st.markdown("---")
            conc=st.text_input(f"Concepto / Destino #{i+1}",
                               placeholder="Excursion o salida / Viaje a...",
                               key=f"{prefix}conc_{i}")
            co1,co2=st.columns(2)
            with co1: hi=st.time_input("Inicio",value=time(0,0),step=300,key=f"{prefix}hi_{i}")
            with co2: hf=st.time_input("Fin",value=time(0,0),step=300,key=f"{prefix}hf_{i}")
            horas_cap.append({"concepto":conc,"hora_ini":hi,"hora_fin":hf})
        ca,cb=st.columns(2)
        with ca:
            if st.button("Anadir horas",type="secondary",use_container_width=True,key=f"{prefix}add_h"):
                st.session_state[n_key]+=1; st.rerun()
        with cb:
            if st.session_state[n_key]>1:
                if st.button("Quitar",type="secondary",use_container_width=True,key=f"{prefix}rm_h"):
                    st.session_state[n_key]-=1; st.rerun()
        st.markdown("---")
        dieta=st.checkbox("🍽️ ¿Dieta hoy?",value=False,key=f"{prefix}dieta")
        dieta_conc=st.text_input(
            "Servicio en horario dieta",
            placeholder="¿Qué servicio da derecho a la dieta? Ej: Excursión Gandía",
            key=f"{prefix}dieta_conc"
        )
        obs=st.text_area("Observaciones (servicios nuevos)",
                         placeholder="Incidencias, servicios nuevos, dudas...",
                         height=65,key=f"{prefix}obs")
    else:
        st.caption("Solo se registrará el nº de servicio para verificación. Sin impacto económico.")

    tipo_val = "extra" if es_extra else "jornada"
    return nombre_sel,fecha_sel,conteo,horas_cap,dieta,dieta_conc,obs,nro_servicio,tipo_val

def vista_conductor():
    st.markdown(CSS,unsafe_allow_html=True)
    if st.session_state.get("envio_ok"):
        info=st.session_state.get("envio_info",{})
        st.markdown(logo_html(),unsafe_allow_html=True); st.divider()
        st.success(f"Registro enviado\n\n{info.get('nombre','')} — {info.get('fecha','')}")
        if st.button("Registrar otro servicio",type="primary",use_container_width=True):
            for k in ["envio_ok","resumen","datos_resumen"]: st.session_state.pop(k,None)
            for k in list(st.session_state.keys()):
                if k.startswith("c_"): del st.session_state[k]
            st.rerun()
        st.markdown('<p class="admin-link">...</p>',unsafe_allow_html=True)
        return
    if st.session_state.get("resumen"):
        datos=st.session_state.get("datos_resumen",{})
        st.markdown(logo_html(),unsafe_allow_html=True)
        st.markdown("<h4 style='text-align:center;color:#1F4E79;'>Revisa tu registro</h4>",unsafe_allow_html=True)
        st.divider()
        st.markdown(f"**{datos.get('nombre','')}** — {datos.get('fecha_str','')}")
        if datos.get("nro_servicio"):
            st.markdown(f"Nº Servicio: **{datos['nro_servicio']}**")
        if datos.get("tipo_registro") == "jornada":
            st.info("📋 Registro de jornada habitual — sin impacto económico")
        else:
            st.markdown(f"Dieta: {'Si' if datos.get('dieta') else 'No'}")
            if datos.get("fijos_ok"):
                st.markdown("Fijos: " + ", ".join(
                    f"{f['tipo']} x{f['cantidad']}" if f['cantidad']>1 else f['tipo']
                    for f in datos["fijos_ok"]))
            if datos.get("horas_ok"):
                st.markdown("Horas: " + ", ".join(
                    f"{h['concepto']} ({h['inicio']}→{h['fin']})" for h in datos["horas_ok"]))
            if datos.get("observaciones"):
                st.markdown(f"Obs: {datos['observaciones']}")
        st.divider()
        c1,c2=st.columns(2)
        with c1:
            if st.button("Corregir",type="secondary",use_container_width=True):
                st.session_state.resumen=False; st.rerun()
        with c2:
            if st.button("Confirmar y Enviar",type="primary",use_container_width=True):
                fila={"conductor":datos["nombre"],"fecha":datos["fecha_iso"],
                      "dieta":datos["dieta"],
                      "servicios_fijos":json.dumps(datos["fijos_ok"],ensure_ascii=False),
                      "servicios_horas":json.dumps(datos["horas_ok"],ensure_ascii=False),
                      "observaciones":datos.get("observaciones") or None,
                      "hoja_servicio":datos.get("nro_servicio") or None,
                      "tipo_registro":datos.get("tipo_registro","extra")}
                with st.spinner("Enviando..."): ok=guardar_fila(fila)
                if ok:
                    st.session_state.envio_ok=True
                    st.session_state.envio_info={"nombre":datos["nombre"],"fecha":datos["fecha_str"]}
                    st.session_state.resumen=False
                    for k in list(st.session_state.keys()):
                        if k.startswith("c_"): del st.session_state[k]
                    st.rerun()
        return
    st.markdown(logo_html(),unsafe_allow_html=True)
    st.markdown("<h4 style='text-align:center;color:#1F4E79;margin:0;'>Registro de Servicios</h4>",unsafe_allow_html=True)
    st.divider()
    nombre_sel,fecha_sel,conteo,horas_cap,dieta_sel,dieta_conc_sel,obs,nro_servicio_sel,tipo_sel = formulario_conductor("c_")
    st.divider()
    if st.button("VER RESUMEN Y CONFIRMAR",type="primary",use_container_width=True):
        errores=[]
        if nombre_sel == "✏️ Mi nombre no aparece en la lista":
            nombre_final = nombre_manual.strip() if nombre_manual.strip() else ""
            if not nombre_final:
                errores.append("Escribe tu nombre completo en el campo de texto.")
        elif nombre_sel == "— Selecciona —":
            nombre_final = ""
            errores.append("Debes seleccionar tu nombre.")
        else:
            nombre_final = extraer_nombre(nombre_sel)

        if tipo_sel == "jornada":
            # Solo necesita nº de servicio
            if not nro_servicio_sel:
                errores.append("Para registro de jornada debes indicar el Nº Servicio.")
            if not errores:
                st.session_state.resumen=True
                st.session_state.datos_resumen={
                    "nombre":nombre_final,"fecha_str":fecha_sel.strftime("%d/%m/%Y"),
                    "fecha_iso":fecha_sel.strftime("%Y-%m-%d"),
                    "dieta":False,"fijos_ok":[],"horas_ok":[],"observaciones":"",
                    "nro_servicio":int(nro_servicio_sel),"tipo_registro":"jornada"}
                st.rerun()
        else:
            fijos_ok=[{"tipo":s,"cantidad":int(conteo[s])} for s in SERVICIOS_GRID if conteo.get(s,0)>0]
            horas_ok=[]
            for i,h in enumerate(horas_cap):
                if h["concepto"].strip():
                    ini=h["hora_ini"]; fin=h["hora_fin"]
                    if ini==time(0,0) and fin==time(0,0):
                        horas_ok.append({"concepto":h["concepto"].strip(),"inicio":"","fin":""})
                    elif fin<=ini: errores.append(f"Horas #{i+1}: fin debe ser posterior al inicio.")
                    else: horas_ok.append({"concepto":h["concepto"].strip(),
                                           "inicio":ini.strftime("%H:%M"),"fin":fin.strftime("%H:%M")})
            if not fijos_ok and not horas_ok and not errores:
                errores.append("Aniade al menos un servicio.")
            if not errores:
                obs_final = obs.strip()
                if dieta_sel and dieta_conc_sel.strip():
                    obs_final = f"[Dieta: {dieta_conc_sel.strip()}] {obs_final}".strip()
                st.session_state.resumen=True
                st.session_state.datos_resumen={
                    "nombre":nombre_final,"fecha_str":fecha_sel.strftime("%d/%m/%Y"),
                    "fecha_iso":fecha_sel.strftime("%Y-%m-%d"),"dieta":dieta_sel,
                    "fijos_ok":fijos_ok,"horas_ok":horas_ok,"observaciones":obs_final,
                    "nro_servicio":int(nro_servicio_sel),"tipo_registro":"extra"}
                st.rerun()
        if errores:
            for msg in errores: st.error(msg)
    st.markdown('<p class="admin-link"><a href="?admin=1" style="color:#aaa;text-decoration:none;">...</a></p>',
                unsafe_allow_html=True)

def vista_admin():
    st.markdown(CSS_ADMIN,unsafe_allow_html=True)
    if "admin_ok" not in st.session_state: st.session_state.admin_ok=False
    if not st.session_state.admin_ok:
        st.title("Administracion — Autocares Alegre"); st.divider()
        pw=st.text_input("Contrasena",type="password")
        if st.button("Entrar",type="primary",use_container_width=True):
            if pw==PASSWORD_ADMIN: st.session_state.admin_ok=True; st.rerun()
            else: st.error("Contrasena incorrecta.")
        return
    c1,c2=st.columns([5,1])
    with c1: st.title("Administracion — Autocares Alegre")
    with c2:
        st.markdown("<br>",unsafe_allow_html=True)
        if st.button("Salir"): st.session_state.admin_ok=False; st.rerun()
    st.divider()

    # ── FILTRO DE PERIODO ────────────────────────────────
    st.markdown("### 📅 Periodo de revisión")
    ini_def, fin_def = periodo_actual()
    col_f1, col_f2, col_f3 = st.columns([2,2,3])
    with col_f1:
        fecha_desde = st.date_input("Desde", value=ini_def,
                                    format="DD/MM/YYYY", key="filtro_desde")
    with col_f2:
        fecha_hasta = st.date_input("Hasta", value=fin_def,
                                    format="DD/MM/YYYY", key="filtro_hasta")
    with col_f3:
        st.markdown("<br>", unsafe_allow_html=True)
        st.caption(f"Periodo activo: **{fecha_desde.strftime('%d/%m/%Y')}** → **{fecha_hasta.strftime('%d/%m/%Y')}**")
    st.divider()
    st.markdown("## Registros enviados")
    if st.button("Actualizar",type="secondary"):
        cargar_datos_brutos.clear(); st.rerun()
    df_brutos_raw = cargar_datos_brutos()
    # Filtrar por el periodo seleccionado
    if not df_brutos_raw.empty:
        df_brutos_raw["fecha"] = pd.to_datetime(df_brutos_raw["fecha"], errors="coerce").dt.date
        df_brutos = df_brutos_raw[
            (df_brutos_raw["fecha"] >= fecha_desde) &
            (df_brutos_raw["fecha"] <= fecha_hasta)
        ].copy()
    else:
        df_brutos = df_brutos_raw
    if df_brutos.empty:
        st.info("Todavia no hay registros.")
    else:
        # Tabla con checkbox para eliminar registros
        filas_v=[]
        for _,row in df_brutos.iterrows():
            revision=necesita_revision(row)
            nro=row.get("hoja_servicio")
            tipo_r=str(row.get("tipo_registro","extra") or "extra")
            filas_v.append({
                "🗑️":False,
                "_id":int(row.get("id",0)),
                "Tipo":"📋 Jornada" if tipo_r=="jornada" else "➕ Extra",
                "Nº Serv.":int(nro) if nro and str(nro) not in ("","nan","None") else "",
                "Conductor":row.get("conductor",""),
                "Fecha":row.get("fecha",""),
                "Dieta":"Si" if str(row.get("dieta","")).lower() in ["true","1"] else "",
                "Servicios Fijos":fijos_a_texto(row.get("servicios_fijos","")),
                "Por Horas":horas_a_texto(row.get("servicios_horas","")),
                "Observaciones":str(row.get("observaciones","")) if row.get("observaciones") else "",
                "Revisar":"REVISAR" if (revision and tipo_r!="jornada") else "",
            })
        df_v=pd.DataFrame(filas_v)
        df_editado_v = st.data_editor(
            df_v, use_container_width=True, hide_index=True,
            column_config={
                "🗑️": st.column_config.CheckboxColumn("🗑️ Borrar", default=False, width="small"),
                "_id": None,  # Ocultar columna ID
                "Tipo": st.column_config.TextColumn("Tipo", width="small"),
            },
            disabled=["Tipo","Conductor","Fecha","Dieta","Servicios Fijos","Por Horas","Observaciones","Revisar","Nº Serv."],
            key="tabla_registros"
        )
        ids_a_borrar = df_editado_v[df_editado_v["🗑️"]==True]["_id"].tolist()
        if ids_a_borrar:
            st.warning(f"Has marcado {len(ids_a_borrar)} registro(s) para eliminar.")
            if st.button(f"🗑️ ELIMINAR {len(ids_a_borrar)} REGISTRO(S) SELECCIONADO(S)",
                         type="primary", use_container_width=True):
                if eliminar_por_ids(ids_a_borrar):
                    st.success(f"✅ {len(ids_a_borrar)} registro(s) eliminado(s).")
                    st.rerun()
    st.divider()
    st.markdown("## Importes calculados")
    st.caption("Haz clic en cualquier celda de importe para corregirla directamente.")
    if not df_brutos.empty:
        df_cierre = calcular_liquidacion(df_brutos)
        if not df_cierre.empty:
            # Recalcular si cambian los datos, pero conservar ediciones manuales
            clave_datos = str(len(df_brutos))
            if (st.session_state.get("df_cierre_clave") != clave_datos or
                    "df_cierre_editado" not in st.session_state):
                st.session_state.df_cierre_editado = df_cierre.copy()
                st.session_state.df_cierre_clave   = clave_datos
            if st.button("↺ Recalcular (descartar correcciones manuales)", type="secondary"):
                st.session_state.df_cierre_editado = df_cierre.copy()
                st.session_state.df_cierre_clave   = clave_datos
                st.rerun()

            # ── Tabla editable ──────────────────────────────
            cols_ed = ["Nº Serv.","Conductor","Fecha","Servicios Fijos","Por Horas",
                       "Observaciones",
                       "IMP. FIJOS","IMP. HORAS","IMP. DIETAS","TOTAL","REVISAR"]
            # Convertir fecha a DD/MM/YYYY para visualización
            df_para_editor = st.session_state.df_cierre_editado.copy()
            if "Fecha" in df_para_editor.columns:
                df_para_editor["Fecha"] = df_para_editor["Fecha"].apply(
                    lambda x: x.strftime("%d/%m/%Y") if hasattr(x,"strftime") else str(x)
                )
            df_ed = st.data_editor(
                df_para_editor[[c for c in cols_ed
                                                    if c in st.session_state.df_cierre_editado.columns]],
                use_container_width=True, hide_index=True,
                column_config={
                    "Nº Serv.":       st.column_config.NumberColumn("Nº Serv.", disabled=True, width="small"),
                    "Conductor":      st.column_config.TextColumn(disabled=True, width="medium"),
                    "Fecha":          st.column_config.TextColumn("Fecha", disabled=True, width="small"),
                    "Servicios Fijos":st.column_config.TextColumn(disabled=True, width="large"),
                    "Por Horas":      st.column_config.TextColumn(disabled=True, width="large"),
                    "IMP. FIJOS":     st.column_config.NumberColumn("FIJOS €",  format="%.2f", step=0.01),
                    "IMP. HORAS":     st.column_config.NumberColumn("HORAS €",  format="%.2f", step=0.01),
                    "IMP. DIETAS":    st.column_config.NumberColumn("DIETAS €", format="%.2f", step=0.01),
                    "TOTAL":          st.column_config.NumberColumn("TOTAL €",  format="%.2f", step=0.01),
                    "Observaciones":  st.column_config.TextColumn("Observaciones / Dieta", disabled=True, width="large"),
                    "REVISAR":        st.column_config.TextColumn("⚠️ Revisar (borra si ya revisado)", width="medium"),
                },
                key="editor_importes"
            )
            # Recalcular TOTAL automáticamente con los valores editados
            df_ed["TOTAL"] = (df_ed["IMP. FIJOS"] + df_ed["IMP. HORAS"] + df_ed["IMP. DIETAS"]).round(2)

            # Solo actualizar las columnas de importe y REVISAR (no tocar Fecha ni columnas de texto)
            cols_actualizar = ["IMP. FIJOS","IMP. HORAS","IMP. DIETAS","TOTAL","REVISAR"]
            for col in cols_actualizar:
                if col in df_ed.columns and col in st.session_state.df_cierre_editado.columns:
                    st.session_state.df_cierre_editado[col] = df_ed[col].values

            # Botón explícito de guardar con confirmación
            col_g1, col_g2 = st.columns([2,3])
            with col_g1:
                if st.button("💾 Guardar correcciones", type="primary"):
                    st.session_state.df_cierre_guardado = st.session_state.df_cierre_editado.copy()
                    st.success("✅ Correcciones guardadas. El Excel y PDF usarán estos importes.")
            with col_g2:
                if st.session_state.get("df_cierre_guardado") is not None:
                    st.info("✅ Tienes correcciones guardadas listas para exportar.")

            # ── Totales ─────────────────────────────────────
            m1,m2,m3,m4=st.columns(4)
            m1.metric("Fijos",  f"{df_ed['IMP. FIJOS'].sum():,.2f} €")
            m2.metric("Horas",  f"{df_ed['IMP. HORAS'].sum():,.2f} €")
            m3.metric("Dietas", f"{df_ed['IMP. DIETAS'].sum():,.2f} €")
            m4.metric("TOTAL",  f"{df_ed['TOTAL'].sum():,.2f} €")
            st.markdown("---")

            # ── Filtro de conductores para el PDF ───────────
            todos_c = sorted(df_ed["Conductor"].unique().tolist())
            sel_pdf  = st.multiselect(
                "📄 Seleccionar conductores para el PDF (vacío = todos):",
                options=todos_c,
                default=[],
                placeholder="Dejar vacío para incluir todos los conductores"
            )
            df_base_export = st.session_state.get("df_cierre_guardado", st.session_state.df_cierre_editado)
            df_para_pdf = df_base_export[df_base_export["Conductor"].isin(sel_pdf)] if sel_pdf else df_base_export

            # ── Botones descarga ─────────────────────────────
            nombre_arch = f"autocares_{date.today().strftime('%Y_%m_%d')}"
            c_xl,c_pdf=st.columns(2)
            with c_xl:
                st.download_button("⬇️ Excel",
                    data=generar_excel_bytes(
                        st.session_state.get("df_cierre_guardado",
                        st.session_state.df_cierre_editado)),
                    file_name=f"{nombre_arch}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary", use_container_width=True)
            with c_pdf:
                if not df_para_pdf.empty:
                    st.download_button("📄 PDF",
                        data=generar_pdf_bytes(df_para_pdf),
                        file_name=f"{nombre_arch}.pdf",
                        mime="application/pdf",
                        type="primary", use_container_width=True)
                else:
                    st.warning("Selecciona al menos un conductor para el PDF.")
    st.divider()
    with st.expander("Anadir Registro Manual"):
        nombre_m,fecha_m,conteo_m,horas_m,dieta_m,dieta_conc_m,obs_m,nro_servicio_m,tipo_m = formulario_conductor("m_")
        if st.button("Guardar Registro Manual",type="primary",use_container_width=True,key="btn_manual"):
            if nombre_m == "— Selecciona —":
                st.error("Selecciona un conductor.")
            else:
                # Extraer nombre sin DNI si viene del desplegable
                if nombre_m == "✏️ Mi nombre no aparece en la lista":
                    nombre_m_final = st.session_state.get("m_nombre_manual","").strip()
                    if not nombre_m_final:
                        st.error("Escribe el nombre en el campo de texto.")
                        nombre_m_final = None
                else:
                    nombre_m_final = extraer_nombre(nombre_m)
                if nombre_m_final:
                  fijos_m=[{"tipo":s,"cantidad":int(conteo_m[s])} for s in SERVICIOS_GRID if conteo_m.get(s,0)>0]
                horas_ok_m=[]
                for h in horas_m:
                    if h["concepto"].strip():
                        ini=h["hora_ini"]; fin=h["hora_fin"]
                        if ini==time(0,0) and fin==time(0,0):
                            horas_ok_m.append({"concepto":h["concepto"].strip(),"inicio":"","fin":""})
                        elif fin>ini:
                            horas_ok_m.append({"concepto":h["concepto"].strip(),
                                               "inicio":ini.strftime("%H:%M"),"fin":fin.strftime("%H:%M")})
                    fila={"conductor":nombre_m_final,"fecha":fecha_m.strftime("%Y-%m-%d"),"dieta":dieta_m,
                          "servicios_fijos":json.dumps(fijos_m,ensure_ascii=False),
                          "servicios_horas":json.dumps(horas_ok_m,ensure_ascii=False),
                          "observaciones":(f"[Dieta: {dieta_conc_m.strip()}] {obs_m}".strip()
                                         if dieta_m and dieta_conc_m.strip()
                                         else obs_m.strip() or None),
                          "hoja_servicio":int(nro_servicio_m) if nro_servicio_m else None,
                          "tipo_registro":tipo_m}
                    if guardar_fila(fila):
                        st.success("Guardado correctamente.")
                        for k in list(st.session_state.keys()):
                            if k.startswith("m_"):
                                del st.session_state[k]
                        st.rerun()
    st.divider()
    with st.expander("📥 Importar desde Excel"):
        st.markdown("Sube tu Excel con las hojas **conducores** y **servicios**.")
        st.caption("Los servicios fijos se calculan solos. Las horas van a Observaciones con [HORAS: X] para que las revises.")
        archivo = st.file_uploader("Selecciona el archivo Excel", type=["xlsx","xls"], key="upload_import")
        if archivo:
            try:
                xl   = pd.ExcelFile(archivo)
                hojas = xl.sheet_names
                st.success(f"Archivo cargado. Hojas encontradas: **{', '.join(hojas)}**")
                col_imp1, col_imp2 = st.columns(2)

                # ── Importar conductores ────────────────────
                with col_imp1:
                    hoja_c = next((h for h in hojas if "conduc" in h.lower()), None)
                    if hoja_c:
                        df_c = pd.read_excel(archivo, sheet_name=hoja_c)
                        st.markdown(f"**👥 Conductores** ({len(df_c)} encontrados)")
                        st.dataframe(df_c.head(5), use_container_width=True, hide_index=True)
                        if st.button("📥 Importar conductores", type="primary", key="btn_imp_cond"):
                            with st.spinner("Importando..."):
                                n, om, errs = importar_conductores_excel(df_c)
                            st.success(f"✅ {n} nuevos · {om} ya existían")
                            if errs: st.warning("Errores: " + str(errs[:3]))
                    else:
                        st.info("No se encontró hoja de conductores.")

                # ── Importar servicios ──────────────────────
                with col_imp2:
                    hoja_s = next((h for h in hojas if "servic" in h.lower()), None)
                    if hoja_s:
                        df_s = pd.read_excel(archivo, sheet_name=hoja_s)
                        st.markdown(f"**📋 Servicios** ({len(df_s)} registros)")
                        st.dataframe(df_s.iloc[:5, [2,4,8,17]], use_container_width=True, hide_index=True)
                        if st.button("📥 Importar servicios", type="primary", key="btn_imp_serv"):
                            with st.spinner("Importando..."):
                                n, errs = importar_servicios_excel(df_s)
                            st.success(f"✅ {n} registros importados")
                            if errs:
                                st.warning(f"{len(errs)} errores")
                                with st.expander("Ver errores"):
                                    for e in errs[:10]: st.text(e)
                    else:
                        st.info("No se encontró hoja de servicios.")
            except Exception as e:
                st.error(f"Error al leer el archivo: {e}")
    st.divider()

    with st.expander("Edicion avanzada"):
        st.caption("Para corregir o borrar registros.")
        if not df_brutos.empty:
            df_show=df_brutos.copy()
            df_show["dieta"]=df_show["dieta"].fillna(False).apply(
                lambda x: bool(x) if isinstance(x,bool) else str(x).lower() in ["true","1"])
            df_show["servicios_fijos"]=df_show["servicios_fijos"].apply(fijos_a_texto)
            df_show["servicios_horas"]=df_show["servicios_horas"].apply(horas_a_texto)
            for col in ["conductor","fecha","observaciones"]:
                df_show[col]=df_show[col].fillna("").astype(str)
            df_show=df_show.rename(columns={
                "conductor":"Conductor","fecha":"Fecha","dieta":"Dieta",
                "servicios_fijos":"Servicios Fijos","servicios_horas":"Por Horas",
                "observaciones":"Observaciones"})
            df_edit=st.data_editor(df_show,use_container_width=True,num_rows="dynamic",hide_index=True,
                column_config={
                    "Conductor":st.column_config.TextColumn(width="medium"),
                    "Fecha":st.column_config.TextColumn(),
                    "Dieta":st.column_config.CheckboxColumn(default=False),
                    "Servicios Fijos":st.column_config.TextColumn(width="large",disabled=True),
                    "Por Horas":st.column_config.TextColumn(width="large",disabled=True),
                    "Observaciones":st.column_config.TextColumn(width="medium"),
                },key="edit_audit")
            if st.button("Guardar cambios",type="secondary",key="btn_audit"):
                df_orig=df_brutos.copy()
                if len(df_edit)==len(df_orig):
                    df_orig["conductor"]=df_edit["Conductor"].values
                    df_orig["fecha"]=df_edit["Fecha"].values
                    df_orig["dieta"]=df_edit["Dieta"].values
                    df_orig["observaciones"]=df_edit["Observaciones"].values
                    if guardar_tabla_completa(df_orig): st.success("Guardado.")

def main():
    if "vista_actual" not in st.session_state: st.session_state.vista_actual="conductor"
    if st.query_params.get("admin")=="1" and st.session_state.vista_actual=="conductor":
        st.session_state.vista_actual="admin"
    if st.session_state.vista_actual=="admin":
        with st.sidebar:
            st.markdown("## Autocares Alegre"); st.markdown("---")
            if st.button("Vista Conductor",use_container_width=True):
                st.session_state.vista_actual="conductor"; st.query_params.clear(); st.rerun()
            st.caption("v9.0")
        vista_admin()
    else:
        vista_conductor()

if __name__=="__main__":
    main()
